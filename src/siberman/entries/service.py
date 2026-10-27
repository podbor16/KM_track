"""
Заявки Siberman: приём (вебхук Тильды, спул, CSV), вкладка «Заявки» админки (статусы, комментарий, журнал,
Excel) и публичный список для iframe в Тильде (только ФИ, формат, статус).
"""
import datetime
import io
import threading
import time
from contextlib import contextmanager
from zoneinfo import ZoneInfo

import mysql.connector
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.config import settings
from src.siberman.db import get_siberman_connection
from src.siberman.entries import db as edb
from src.siberman.entries.parser import BadEntry, is_test_ping, parse_csv, parse_entry

KRS = ZoneInfo("Asia/Krasnoyarsk")
STATUSES = ("review", "admitted", "rejected", "withdrawn")
STATUS_LABELS = {"review": "⌛️ на изучении", "admitted": "✅ допущена", "rejected": "❌ не соответствует",
                 "withdrawn": "🚫 отказ от жеребьёвки"}
FORMAT_LABELS = {"individual": "Индивидуальное участие", "relay": "Эстафетная команда"}
PUBLIC_TTL_S = 5.0
_public_cache: dict = {}
_lock = threading.Lock()


class EntryError(Exception):
    """Ошибка, которую директор видит как есть (409)."""


class EntriesUnavailable(EntryError):
    """БД недоступна (503)."""


def clear_cache() -> None:
    with _lock:
        _public_cache.clear()


@contextmanager
def _db():
    """Соединение с БД заявок; недоступна или оборвалась посреди запроса — EntriesUnavailable (503)."""
    conn = get_siberman_connection()
    if conn is None:
        raise EntriesUnavailable("база данных временно недоступна")
    try:
        yield conn
    except mysql.connector.Error as e:
        raise EntriesUnavailable("база данных временно недоступна") from e
    finally:
        conn.close()


def now_krs() -> datetime.datetime:
    return datetime.datetime.now(KRS).replace(tzinfo=None, microsecond=0)


def _fio(e: dict) -> str:
    return f"{e['surname']} {e['name']}".strip()


def _norm(s: str) -> str:
    return (s or "").lower().replace("ё", "е")


# ── Приём ───────────────────────────────────────────────────────────────────

def ingest(year: int, body: dict, received_at: datetime.datetime):
    """Тело вебхука → заявка. None — тестовый запрос Тильды; True/False — новая / уже была.
    BadEntry — не заявка (в failed/), прочие исключения — сбой БД (повтор из спула)."""
    if is_test_ping(body):
        return None
    e = parse_entry(body, received_at.date())
    with _db() as conn:
        new = edb.upsert(conn, year, e, received_at)
    clear_cache()
    return new


def import_csv(year: int, data: bytes) -> dict:
    """Разовая загрузка выгрузки Тильды: новые добавляются (время — Date из CSV), известные обновляют поля формы."""
    try:
        rows, errors = parse_csv(data, now_krs().date())
    except BadEntry as e:
        raise EntryError(str(e)) from None
    with _db() as conn:
        seen = edb.known_tranids(conn, [e["tranid"] for e, _ in rows])
        total = len(rows) + len(errors)
        known = {t for t, y in seen.items() if y == year}
        to_write = []
        for e, at in rows:
            if seen.get(e["tranid"], year) != year:
                errors.append(f"{e['surname']} {e['name']}: заявка уже учтена в {seen[e['tranid']]} году")
            else:
                to_write.append((e, at))
        for e, at in to_write:
            edb.upsert(conn, year, e, at or now_krs())
    clear_cache()
    return {"rows": total, "new": sum(1 for e, _ in to_write if e["tranid"] not in known),
            "known": sum(1 for e, _ in to_write if e["tranid"] in known), "errors": errors}


# ── Админка ─────────────────────────────────────────────────────────────────

def webhook_path(year: int):
    secret = settings.SIBERMAN_WEBHOOK_SECRET
    return f"/webhook/siberman/{secret}/{year}" if secret else None


def admin_list(year: int) -> dict:
    with _db() as conn:
        rows, years = edb.list_year(conn, year), edb.years(conn)
    counts = {"total": len(rows), **{s: sum(1 for r in rows if r["status"] == s) for s in STATUSES}}
    for n, r in enumerate(rows, start=1):
        r["n"], r["comment"] = n, r["comment"] or ""
    return {"year": year, "entries": rows, "counts": counts, "years": sorted(set(years) | {year}),
            "webhook_path": webhook_path(year)}


def _entry(conn, year: int, entry_id: int) -> dict:
    e = edb.get(conn, entry_id)
    if e is None or e["race_year"] != year:
        raise EntryError("нет такой заявки")
    return e


def entry_card(year: int, entry_id: int) -> dict:
    with _db() as conn:
        e = _entry(conn, year, entry_id)
        e["audit"] = edb.audit(conn, entry_id)
    return e


def set_status(year: int, entry_id: int, status: str, user: str) -> dict:
    if status not in STATUSES:
        raise EntryError("нет такого статуса")
    with _db() as conn:
        e = _entry(conn, year, entry_id)
        if e["status"] != status:
            edb.update(conn, entry_id, "status", e["status"], status, user)
    clear_cache()
    return {"ok": True}


def set_comment(year: int, entry_id: int, comment: str, user: str) -> dict:
    comment = (comment or "").strip()[:2000]
    with _db() as conn:
        e = _entry(conn, year, entry_id)
        if (e["comment"] or "") != comment:
            edb.update(conn, entry_id, "comment", e["comment"] or "", comment, user)
    return {"ok": True}


EXPORT_COLUMNS = (("n", "№", 5), ("submitted_at", "Дата заявки", 18), ("fio", "Фамилия Имя", 28),
                  ("format", "Формат участия", 22), ("gender", "Пол", 6), ("birthdate", "Дата рождения", 13),
                  ("city", "Город", 18), ("email", "E-mail", 28), ("phone", "Телефон", 18),
                  ("status", "Статус", 22), ("comment", "Комментарий", 30), ("team_name", "Команда", 20),
                  ("relay_swim", "Плавание", 22), ("relay_bike", "Велосипед", 22), ("relay_run", "Бег", 22),
                  ("wildcard_reason", "Почему wild card", 30), ("results_text", "О результатах", 30),
                  ("result_url", "Ссылка на результаты", 30))


def export_xlsx(year: int) -> bytes:
    rows = admin_list(year)["entries"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"Заявки {year}"
    ws.append([t for _, t, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        v = {**r, "fio": _fio(r), "format": FORMAT_LABELS[r["format"]], "status": STATUS_LABELS[r["status"]],
             "gender": {"M": "М", "F": "Ж"}.get(r["gender"], "")}
        ws.append([v.get(k) if v.get(k) is not None else "" for k, _, _ in EXPORT_COLUMNS])
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":         # текст из формы, начинающийся с «=», — не формула
                cell.data_type = "s"
    for i, (key, _, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        fmt = {"submitted_at": "DD.MM.YYYY HH:MM", "birthdate": "DD.MM.YYYY"}.get(key)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Публичный список ────────────────────────────────────────────────────────

def public_list(year: int) -> dict:
    """По алфавиту (фамилия, имя; ё = е); только ФИ, формат, статус. Кэш воркера на PUBLIC_TTL_S."""
    with _lock:
        hit = _public_cache.get(year)
        if hit and time.monotonic() - hit[0] < PUBLIC_TTL_S:
            return hit[1]
    with _db() as conn:
        rows = edb.public_rows(conn, year)
    rows.sort(key=lambda r: (_norm(r["surname"]), _norm(r["name"]), r["id"]))
    data = {"year": year, "total": len(rows), "admitted": sum(1 for r in rows if r["status"] == "admitted"),
            "entries": [{"n": n, "fio": _fio(r), "format": FORMAT_LABELS[r["format"]], "status": r["status"]}
                        for n, r in enumerate(rows, start=1)]}
    with _lock:
        _public_cache[year] = (time.monotonic(), data)
    return data
