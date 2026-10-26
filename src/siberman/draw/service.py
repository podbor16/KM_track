"""
Жеребьёвка Siberman: снимки состояния (админка/публика), операции записи, кэш, фиды vMix.

Каждая запись — в транзакции под блокировкой строки draw_config года (SELECT … FOR UPDATE):
операции года идут строго по очереди, двойной клик не создаст дубль. Каждая запись
увеличивает draw_config.version; публичный снимок пересобирается при смене версии,
а сама версия читается из БД не чаще раза в CACHE_TTL_S на воркер.
"""
import csv
import datetime
import hashlib
import hmac
import io
import logging
import threading
import time

import mysql.connector
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import settings
from src.siberman.db import get_siberman_connection
from src.siberman.draw import db as ddb
from src.siberman.draw.importer import contact_for, diff_applicants, import_report, parse_applicants, parse_contacts
from src.siberman.draw.logic import (
    CATEGORIES, CATEGORY_LABELS, ROUND_LABELS, build_lists, category_of, check_pick,
    in_round_pool, placement, round_counters,
)
from src.siberman.draw.payments import PAYMENT_ICONS, auto_invite_allowed, next_invitee
from src.siberman.draw.simulation import win_probabilities
from src.siberman.hall import service as hall_service
from src.siberman.hall.service import HallUnavailable
from src.siberman.persons.linking import plain_key


log = logging.getLogger(__name__)


class DrawError(Exception):
    """Ошибка, которую оператор видит как есть."""


class DrawUnavailable(DrawError):
    """БД недоступна: роутер отвечает 503 (а не 409), vMix/страница повторят запрос."""


CACHE_TTL_S = 1.0
# 2000 прогонов: на VPS (3 ядра EPYC) ~0,9 с вместо ~2,2 с при 5000 — пересчёт идёт в каждом
# из 3 воркеров и под GIL тормозит их ответы; погрешность ±1 п.п. (замер 2026-10-14).
SIMS = 2000
FEEDS = ("men", "women", "relay", "wildcard", "last")
FEED_COLUMNS = {
    "men": ("num", "name", "city", "section"),          # оплаты в эфире нет: счета — после трансляции
    "women": ("num", "name", "city", "section"),
    "relay": ("num", "name", "city", "section"),
    "wildcard": ("num", "name", "city", "category"),
    "last": ("name", "city", "category", "num", "round"),
}
SECTION_LABELS = {"winner": "Победитель", "waitlist": "Лист ожидания", None: ""}
AQUARIUM_GROUPS = {"relay": ("Эстафеты", "blue"), "repeat": ("Повторные", "black"),
                   "novice": ("Новички", "black"), "individual": ("Личники", "black")}

_cache: dict[int, tuple[int, float, dict]] = {}
_state_cache: dict[int, tuple[tuple, dict]] = {}  # готовый public_state: (ключ версий, ответ)
_host_cache: dict[int, tuple[int, dict]] = {}     # готовый host_state: (версия снимка, ответ)
_probs: dict[int, tuple[int, dict | None]] = {}
_probs_running: set[int] = set()
_lock = threading.Lock()
_year_locks: dict[int, threading.Lock] = {}


def clear_cache() -> None:
    with _lock:
        _cache.clear()
        _state_cache.clear()
        _host_cache.clear()
        _probs.clear()


def _connect():
    conn = get_siberman_connection()
    if conn is None:
        raise DrawUnavailable("БД недоступна")
    return conn


def _year_lock(year: int) -> threading.Lock:
    with _lock:
        return _year_locks.setdefault(year, threading.Lock())


def _empty_config() -> dict:
    cfg = {f: None for f in ddb.CONFIG_INT_FIELDS}
    return {**cfg, "current_round": 0, "version": 0, "auto_invite_until": None, "imported_at": None}


def _snapshot(conn, year: int, create: bool = True) -> dict:
    """create=False — путь публичного чтения: без INSERT, для года без строки — пустое состояние."""
    if create:
        config = ddb.load_config(conn, year)
    else:
        config = ddb.get_config(conn, year)
        if config is None:
            return {"year": year, "config": _empty_config(), "applicants": [], "picks": [],
                    "lists": build_lists({}, [], _empty_config())}
    apps = ddb.load_applicants(conn, year)
    picks = ddb.load_picks(conn, year)
    lists = build_lists({a["id"]: a for a in apps}, picks, config)
    return {"year": year, "config": config, "applicants": apps, "picks": picks, "lists": lists}


def _applicant(snap: dict, applicant_id: int) -> dict:
    a = next((x for x in snap["applicants"] if x["id"] == applicant_id), None)
    if a is None:
        raise DrawError("нет такой заявки")
    return a


def _active_pick(snap: dict, pick_id: int) -> dict:
    p = next((x for x in snap["picks"] if x["id"] == pick_id), None)
    if p is None or p["cancelled_at"] is not None:
        raise DrawError("запись не найдена или уже отменена")
    return p


def _fio(a: dict) -> str:
    return f"{a['surname']} {a['name']}".strip()


def _write(year: int, op):
    """op(conn, snap) в транзакции под блокировкой строки года; version += 1."""
    conn = _connect()
    try:
        ddb.ensure_config(conn, year)
        conn.start_transaction()
        ddb.lock_config(conn, year)
        result = op(conn, _snapshot(conn, year))
        ddb.bump_version(conn, year)
        conn.commit()
        return result
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


# ── Записи ──────────────────────────────────────────────────────────────────

def pick(year: int, applicant_id: int, user: str = "") -> dict:
    def op(conn, snap):
        a = _applicant(snap, applicant_id)
        rnd = snap["config"]["current_round"]
        reason = check_pick(a, rnd, snap["lists"], snap["config"], {x["id"]: x for x in snap["applicants"]})
        if reason:
            raise DrawError(f"{_fio(a)}: {reason}")
        category = category_of(a)
        pick_id = ddb.insert_pick(conn, year, a["id"], rnd, category)
        return {"pick_id": pick_id, "category": category, "position": len(snap["lists"][category]) + 1}
    return _write(year, op)


def _reset_payment(conn, year: int, snap: dict, applicant_id: int, user: str) -> None:
    """Вытесненный/отменённый участник больше не в списке — его статус оплаты не должен висеть."""
    a = next((x for x in snap["applicants"] if x["id"] == applicant_id), None)
    if a and a["payment_status"] != "none":
        ddb.update_payment_status(conn, applicant_id, "none")
        ddb.audit(conn, year, applicant_id, "payment", a["payment_status"], "none", None, user)


def cancel(year: int, pick_id: int, user: str = "") -> dict:
    def op(conn, snap):
        p = _active_pick(snap, pick_id)
        ddb.cancel_pick(conn, pick_id)
        ddb.audit(conn, year, p["applicant_id"], "cancel_pick", str(pick_id), "", None, user)
        _reset_payment(conn, year, snap, p["applicant_id"], user)
        return {"ok": True}
    return _write(year, op)


def replace(year: int, pick_id: int, applicant_id: int, user: str = "") -> dict:
    """Исправить: другой участник на ту же точку хода жеребьёвки (раунд и очерёдность исходной записи).
    Зачёт может быть другим (ввели женщину вместо вытянутого мужчины): новый встаёт в свой зачёт так,
    будто его ввели сразу, соседи по зачётам сдвигаются. Кто из-за этого перешёл между победителями
    и листом ожидания — в ответе (shifted)."""
    def op(conn, snap):
        p = _active_pick(snap, pick_id)
        a = _applicant(snap, applicant_id)
        reason = check_pick(a, p["round"], snap["lists"])
        if reason:
            raise DrawError(f"{_fio(a)}: {reason}")
        category = category_of(a)
        new_id = ddb.insert_pick(conn, year, a["id"], p["round"], category, position_key=p["position_key"])
        ddb.cancel_pick(conn, pick_id, replaced_by=new_id)
        ddb.audit(conn, year, a["id"], "replace_pick", str(p["applicant_id"]), str(a["id"]), p["applicant_id"], user)
        _reset_payment(conn, year, snap, p["applicant_id"], user)
        picks = [{**x, "cancelled_at": "исправлено"} if x["id"] == pick_id else x for x in snap["picks"]]
        picks.append({"id": new_id, "applicant_id": a["id"], "round": p["round"], "category": category,
                      "position_key": p["position_key"], "created_at": None, "cancelled_at": None})
        after = build_lists({x["id"]: x for x in snap["applicants"]}, picks, snap["config"])
        before = {r["applicant_id"]: r["section"] for rows in snap["lists"].values() for r in rows}
        rows = [(c, r) for c, rows in after.items() for r in rows]
        shifted = [{"surname": r["surname"], "name": r["name"], "category": c, "position": r["position"],
                    "section": r["section"]}
                   for c, r in rows if r["applicant_id"] != a["id"] and before.get(r["applicant_id"]) != r["section"]]
        position = next(r["position"] for c, r in rows if r["applicant_id"] == a["id"])
        return {"pick_id": new_id, "category": category, "position": position, "shifted": shifted}
    return _write(year, op)


def set_round(year: int, rnd: int, user: str = "") -> dict:
    if rnd not in ROUND_LABELS:
        raise DrawError("нет такого раунда")

    def op(conn, snap):
        ddb.set_round(conn, year, rnd)
        return {"round": rnd}
    return _write(year, op)


def save_config(year: int, values: dict, user: str = "") -> dict:
    clean = {}
    for k in ddb.CONFIG_INT_FIELDS:
        if k not in values:
            continue
        v = values[k]
        if v is not None and (not isinstance(v, int) or isinstance(v, bool) or not 0 <= v <= 1000):
            raise DrawError(f"{k}: нужно целое число от 0 до 1000")
        clean[k] = v
    if "auto_invite_until" in values:
        clean["auto_invite_until"] = values["auto_invite_until"]

    def op(conn, snap):
        ddb.save_config(conn, year, clean)
        return {"ok": True}
    return _write(year, op)


def _against(snap: dict, apps: list[dict]) -> dict:
    """Сравнение с загруженной таблицей (организаторы меняют шансы до самой жеребьёвки) и запрет записи."""
    loaded = snap["config"]["imported_at"] is not None
    return {"diff": diff_applicants(snap["applicants"], apps) if loaded else None, "blocked": bool(snap["picks"])}


def import_applicants(year: int, file_bytes: bytes, dry_run: bool = True, user: str = "") -> dict:
    """Таблица заменяет заявки года целиком — можно перезаливать сколько угодно до первого вытянутого.
    Проверка (dry_run) ничего не пишет: отчёт, ошибки и что изменится относительно загруженной таблицы."""
    apps, errors = parse_applicants(file_bytes)
    report = import_report(apps)
    if dry_run or errors:
        conn = _connect()
        try:
            snap = _snapshot(conn, year, create=False)
        finally:
            conn.close()
        return {"report": report, "errors": errors, "written": False, **_against(snap, apps)}

    def op(conn, snap):
        if snap["picks"]:
            raise DrawError("жеребьёвка уже началась — загрузка таблицы запрещена (сначала «Очистить журнал»)")
        result = {"report": report, "errors": [], "written": True, **_against(snap, apps)}
        ddb.replace_applicants(conn, year, apps)
        return result
    return _write(year, op)


def reset(year: int, confirm: str, user: str = "") -> dict:
    if confirm != str(year):
        raise DrawError(f"для очистки введите {year}")

    def op(conn, snap):
        ddb.reset_draw(conn, year)
        return {"ok": True}
    return _write(year, op)


def set_payment(year: int, applicant_id: int, status: str, comment: str = "", user: str = "",
                today: datetime.date | None = None) -> dict:
    if status not in PAYMENT_ICONS:
        raise DrawError("нет такого статуса оплаты")
    today = today or datetime.date.today()

    def op(conn, snap):
        a = _applicant(snap, applicant_id)
        placed = placement(snap["lists"], a["id"])
        if placed is None:
            raise DrawError(f"{_fio(a)}: ещё не вытянут")
        old, old_comment = a["payment_status"], a["payment_comment"] or ""
        ddb.update_payment(conn, a["id"], status, comment)
        if old != status:
            ddb.audit(conn, year, a["id"], "payment", old, status, None, user)
        elif old_comment != comment[:500]:
            ddb.audit(conn, year, a["id"], "payment_comment", old_comment[:200], comment[:200], None, user)
        result = {"auto_invited": None, "waitlist_empty": False, "suggest_revert": None, "already_invited": None}
        if status == "unpaid" and old != "unpaid" and auto_invite_allowed(snap["config"]["auto_invite_until"], today):
            rows = snap["lists"][placed[0]]
            for r in rows:
                if r["applicant_id"] == a["id"]:
                    r["payment_status"] = status
            prev_id = ddb.last_auto_invite(conn, year, a["id"])
            prev = next((x for x in snap["applicants"] if x["id"] == prev_id), None)
            nxt = next_invitee(rows)
            if prev and prev["payment_status"] == "pending":
                # цепочка ❌→✅→❌: приглашённый в прошлый раз ещё ждёт — второго не зовём
                result["already_invited"] = {"applicant_id": prev["id"], "surname": prev["surname"],
                                             "name": prev["name"]}
            elif nxt is None:
                result["waitlist_empty"] = True
            else:
                ddb.update_payment_status(conn, nxt["applicant_id"], "pending")
                ddb.audit(conn, year, nxt["applicant_id"], "auto_invite", "none", "pending", a["id"], user)
                result["auto_invited"] = {"applicant_id": nxt["applicant_id"], "surname": nxt["surname"],
                                          "name": nxt["name"], "category": placed[0], "position": nxt["position"]}
        if old == "unpaid" and status != "unpaid":
            invited_id = ddb.last_auto_invite(conn, year, a["id"])
            invited = next((x for x in snap["applicants"] if x["id"] == invited_id), None)
            if invited and invited["payment_status"] == "pending":
                result["suggest_revert"] = {"applicant_id": invited["id"], "surname": invited["surname"],
                                            "name": invited["name"]}
        return result
    return _write(year, op)


# ── Чтение: админка ─────────────────────────────────────────────────────────

def admin_state(year: int) -> dict:
    conn = _connect()
    try:
        snap = _snapshot(conn, year)
        contacts_count = len(ddb.load_contacts(conn, year))
    finally:
        conn.close()
    cfg, rnd = snap["config"], snap["config"]["current_round"]
    by_id = {a["id"]: a for a in snap["applicants"]}
    positions = {r["pick_id"]: r["position"] for rows in snap["lists"].values() for r in rows}
    for a in snap["applicants"]:
        placed = placement(snap["lists"], a["id"])
        a["placed"] = {"category": placed[0], "position": placed[1]} if placed else None
        a["in_round"] = in_round_pool(a, rnd)
        a["category"] = category_of(a)
        a["block"] = check_pick(a, rnd, snap["lists"], cfg, by_id)
    journal = [{
        "pick_id": p["id"], "applicant_id": p["applicant_id"],
        "surname": by_id[p["applicant_id"]]["surname"], "name": by_id[p["applicant_id"]]["name"],
        "round": p["round"], "category": p["category"], "created_at": p["created_at"],
        "cancelled_at": p["cancelled_at"], "replaced_by": p["replaced_by"], "position": positions.get(p["id"]),
    } for p in sorted(snap["picks"], key=lambda p: p["id"], reverse=True)]
    return {"year": year, "config": cfg, "round_labels": ROUND_LABELS, "category_labels": CATEGORY_LABELS,
            "applicants": snap["applicants"], "lists": snap["lists"], "journal": journal,
            "counters": round_counters(snap["applicants"], snap["lists"], rnd),
            "host_link": f"/draw/{year}/host/{host_token(year)}", "contacts_count": contacts_count}


def applicant_summary(year: int, applicant_id: int) -> dict:
    conn = _connect()
    try:
        snap = _snapshot(conn, year)
        a = _applicant(snap, applicant_id)
        history = ddb.siberman_history(conn, a["surname"], a["name"])
        audit = ddb.load_audit(conn, year, applicant_id)
    finally:
        conn.close()
    placed = placement(snap["lists"], applicant_id)
    where = None
    if placed:
        row = next(r for r in snap["lists"][placed[0]] if r["applicant_id"] == applicant_id)
        where = {"category": placed[0], "position": placed[1], "round": row["round"], "created_at": row["created_at"]}
    return {"applicant": a, "history": history, "placement": where, "audit": audit}


# ── Счета: контакты из заявок Тильды и Excel ───────────────────────────────

PAYMENT_LABELS = {"none": "—", "pending": "⏳ ожидает оплаты", "paid": "✅ оплачено", "unpaid": "❌ не оплачено"}
EXPORT_COLUMNS = (("category", "Зачёт", 12), ("position", "№", 5), ("section", "Раздел", 15),
                  ("fio", "Фамилия Имя", 28), ("city", "Город", 18), ("email", "E-mail", 30), ("phone", "Телефон", 18),
                  ("payment", "Оплата", 18), ("comment", "Комментарий", 30), ("wildcard", "Wild card", 10),
                  ("round", "Раунд", 14))


def import_contacts(year: int, file_bytes: bytes, user: str = "") -> dict:
    """CSV заявок Тильды → e-mail и телефоны (заменяют прошлую загрузку). Отчёт: сколько людей и у каких
    победителей / приглашённых контакта нет."""
    contacts, report = parse_contacts(file_bytes)
    if not contacts:
        raise DrawError("в файле нет заявок — нужна CSV-выгрузка заявок Тильды (столбцы surname, Name, Email, Phone, format)")

    def op(conn, snap):
        ddb.replace_contacts(conn, year, contacts)
        rows = _payment_rows(snap, contacts)
        return {**report, "matched": sum(1 for a in snap["applicants"] if contact_for(a, contacts)),
                "applicants": len(snap["applicants"]),
                "missing": [r["fio"] for r in rows if not r["email"] and not r["phone"]]}
    return _write(year, op)


def _payment_rows(snap: dict, contacts: dict) -> list[dict]:
    """Кому выставлять счёт: победители (с wild card) и приглашённые из листа ожидания (⏳)."""
    by_id = {a["id"]: a for a in snap["applicants"]}
    out = []
    for c in CATEGORIES:
        for r in snap["lists"][c]:
            if r["section"] == "waitlist" and r["payment_status"] != "pending":
                continue
            a = by_id[r["applicant_id"]]
            ct = contact_for(a, contacts) or {"email": "", "phone": ""}
            out.append({"category": CATEGORY_LABELS[c], "position": r["position"],
                        "section": SECTION_LABELS[r["section"]], "fio": _fio(a), "city": a["city"],
                        "email": ct["email"], "phone": ct["phone"], "payment": PAYMENT_LABELS[a["payment_status"]],
                        "comment": a["payment_comment"] or "", "wildcard": "да" if a["is_wildcard"] else "",
                        "round": ROUND_LABELS[r["round"]]})
    return out


def payment_rows(year: int) -> list[dict]:
    conn = _connect()
    try:
        snap = _snapshot(conn, year)
        contacts = ddb.load_contacts(conn, year)
    finally:
        conn.close()
    return _payment_rows(snap, contacts)


def payment_export(year: int) -> bytes:
    """Excel для счетов (Коле): шапка закреплена, фильтры, у кого нет контакта — подсвечено."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Счета"
    ws.append([title for _, title, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    keys = [k for k, _, _ in EXPORT_COLUMNS]
    missing = PatternFill("solid", fgColor="FFF3C4")
    for r in payment_rows(year):
        ws.append([r[k] for k in keys])
        if not r["email"] and not r["phone"]:
            for k in ("email", "phone"):
                ws.cell(row=ws.max_row, column=keys.index(k) + 1).fill = missing
    for i, (_, _, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Бэкап всей жеребьёвки в Excel ───────────────────────────────────────────

AUDIT_LABELS = {"replace_pick": "исправление", "cancel_pick": "отмена", "payment": "оплата",
                "payment_comment": "комментарий к оплате", "auto_invite": "приглашение из листа"}
CONFIG_LABELS = (("capacity_men", "Мест: мужчины"), ("capacity_women", "Мест: женщины"),
                 ("capacity_relay", "Мест: эстафеты"), ("quota_repeat", "Квота повторных"),
                 ("quota_novice", "Квота новичков"), ("auto_invite_until", "Автопередача до"),
                 ("current_round", "Текущий раунд"), ("imported_at", "Таблица загружена"))


def _sheet(wb: Workbook, title: str, columns: tuple, rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append([t for t, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def draw_backup(year: int) -> bytes:
    """Весь ход жеребьёвки: каждая запись по очерёдности вытягивания (исправление — сразу за исправленной,
    отменённые тоже), итоговые списки, журнал действий и параметры — чтобы восстановить картину без БД."""
    conn = _connect()
    try:
        snap = _snapshot(conn, year)
        audit = ddb.load_year_audit(conn, year)
    finally:
        conn.close()
    by_id = {a["id"]: a for a in snap["applicants"]}
    pick_by_id = {p["id"]: p for p in snap["picks"]}
    placed = {r["pick_id"]: r for rows in snap["lists"].values() for r in rows}
    fio_of = lambda aid: _fio(by_id[aid]) if aid in by_id else ""  # noqa: E731

    def status(p):
        if p["replaced_by"]:
            return f"исправлено → {fio_of(pick_by_id[p['replaced_by']]['applicant_id'])}"
        return "отменено" if p["cancelled_at"] else "в списке"

    course = []
    for i, p in enumerate(sorted(snap["picks"], key=lambda p: (p["position_key"], p["id"])), start=1):
        a, row = by_id[p["applicant_id"]], placed.get(p["id"])
        course.append([i, p["created_at"], ROUND_LABELS[p["round"]], _fio(a), a["city"],
                       CATEGORY_LABELS[p["category"]], row["position"] if row else "",
                       SECTION_LABELS[row["section"]] if row else "", "да" if a["is_wildcard"] else "",
                       status(p), p["cancelled_at"] or ""])
    final = [[CATEGORY_LABELS[c], r["position"], SECTION_LABELS[r["section"]], f"{r['surname']} {r['name']}".strip(),
              r["city"], "да" if r["is_wildcard"] else "", ROUND_LABELS[r["round"]], r["created_at"],
              PAYMENT_LABELS[r["payment_status"]], by_id[r["applicant_id"]]["payment_comment"] or ""]
             for c in CATEGORIES for r in snap["lists"][c]]
    log_rows = []
    for x in audit:
        old, new = x["old_value"], x["new_value"]
        if x["action"] == "replace_pick":
            old, new = fio_of(int(old)) if old.isdigit() else old, fio_of(int(new)) if new.isdigit() else new
        elif x["action"] in ("payment", "auto_invite"):
            old, new = PAYMENT_LABELS.get(old, old), PAYMENT_LABELS.get(new, new)
        log_rows.append([x["created_at"], AUDIT_LABELS.get(x["action"], x["action"]), fio_of(x["applicant_id"]),
                         old, new, fio_of(x["ref_applicant_id"]) if x["ref_applicant_id"] else "", x["admin_user"]])
    cfg = snap["config"]
    params = [[label, ROUND_LABELS.get(cfg[k], cfg[k]) if k == "current_round" else cfg.get(k)]
              for k, label in CONFIG_LABELS]
    params.append(["Выгружено", datetime.datetime.now().replace(microsecond=0)])

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Ход жеребьёвки", (("№", 6), ("Время", 20), ("Раунд", 14), ("Фамилия Имя", 28), ("Город", 18),
                                  ("Зачёт", 11), ("Место", 7), ("Раздел", 15), ("Wild card", 10), ("Статус", 34),
                                  ("Отменено / исправлено", 20)), course)
    _sheet(wb, "Итог", (("Зачёт", 11), ("№", 5), ("Раздел", 15), ("Фамилия Имя", 28), ("Город", 18),
                        ("Wild card", 10), ("Раунд", 14), ("Время", 20), ("Оплата", 18), ("Комментарий", 30)), final)
    _sheet(wb, "Журнал действий", (("Время", 20), ("Действие", 22), ("Участник", 28), ("Было", 28),
                                   ("Стало", 28), ("Связан с", 28), ("Кто", 14)), log_rows)
    _sheet(wb, "Параметры", (("Параметр", 22), ("Значение", 22)), params)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, datetime.datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM:SS"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Страница ведущих (секретная ссылка) ──────────────────────────────────────

HOST_FIELDS = ("id", "format", "surname", "name", "gender", "city", "pool", "is_wildcard", "chances",
               "starts_individual", "starts_relay", "draw_accumulated", "is_local", "duathlon_222", "race_24h",
               "tattoo", "relay_swim", "relay_bike", "relay_run", "wildcard_reason", "results_text")


def host_token(year: int) -> str:
    """Ключ ссылки ведущих: HMAC от SECRET_KEY — не хранится в БД, у каждого года свой."""
    return hmac.new(settings.SECRET_KEY.encode(), f"draw-host:{year}".encode(), hashlib.sha256).hexdigest()[:32]


def host_token_ok(year: int, token: str) -> bool:
    return hmac.compare_digest(token or "", host_token(year))


def match_person(a: dict, by_key: dict):
    """Человек реестра по ФИ заявки (в любом порядке); тёзки в реестре или нет такого — None."""
    for key in (plain_key(a["surname"], a["name"]), plain_key(a["name"], a["surname"])):
        if key in by_key:
            return by_key[key]
    return None


def _athlete_brief(person_id) -> dict | None:
    """Со страницы спортсмена: финиши, лучший результат гонки, рейтинг, действующие рекорды."""
    if person_id is None:
        return None
    try:
        page = hall_service.athlete(person_id)
    except HallUnavailable:
        return None
    if not page:
        return None
    best = next((b for b in page["best"] if b["seg"] == "race"), None)
    return {"id": person_id, "slug": page["person"]["slug"], "summary": page["summary"], "best_race": best,
            "best_place": page["best_place"], "rating": page["rating"]}


def _build_host_state(year: int, snap: dict) -> dict:
    by_id = {a["id"]: a for a in snap["applicants"]}
    drawn = [by_id[r["applicant_id"]] for rows in snap["lists"].values() for r in rows]
    conn = _connect()
    try:
        histories = ddb.siberman_histories(conn, drawn)
        persons = ddb.person_ids_by_key(conn)
    finally:
        conn.close()
    people = {a["id"]: {**{k: a[k] for k in HOST_FIELDS}, "history": histories[a["id"]],
                        "athlete": _athlete_brief(match_person(a, persons))} for a in drawn}
    active = [p for p in snap["picks"] if p["cancelled_at"] is None]
    return {
        "year": year, "current_round": snap["config"]["current_round"], "round_labels": ROUND_LABELS,
        "category_labels": CATEGORY_LABELS,
        "lists": {c: [{k: r[k] for k in ("applicant_id", "position", "section", "round", "is_wildcard")}
                      for r in rows] for c, rows in snap["lists"].items()},
        "people": people,
        "last": max(active, key=lambda p: p["id"])["applicant_id"] if active else None,
    }


def host_state(year: int) -> dict:
    """Страница ведущих: зачёты, карточки всех вытянутых (заявка, история Siberman, страница спортсмена),
    последний вытянутый. Пересчитывается только при изменении жеребьёвки (версия снимка)."""
    snap = _cached_snapshot(year)
    hit = _host_cache.get(year)
    if hit and hit[0] == snap["version"]:
        return hit[1]
    state = _build_host_state(year, snap)
    with _lock:
        _host_cache[year] = (snap["version"], state)
    return state


# ── Чтение: публика (кэш) ───────────────────────────────────────────────────

def _fresh(hit, now: float) -> bool:
    return bool(hit) and now - hit[1] < CACHE_TTL_S


def _store(year: int, version: int, snap: dict) -> None:
    """Не затираем более новую версию более старой (запись из соседнего потока/воркера)."""
    with _lock:
        cur = _cache.get(year)
        if cur is None or cur[0] <= version:
            _cache[year] = (version, time.monotonic(), snap)


def _cached_snapshot(year: int) -> dict:
    hit = _cache.get(year)
    if _fresh(hit, time.monotonic()):
        return hit[2]
    lock = _year_lock(year)  # single-flight: на промах кэша — один запрос в БД на воркер
    if not lock.acquire(blocking=False):
        if hit:  # кто-то уже обновляет — не ждём, отдаём прошлый снимок
            return hit[2]
        lock.acquire()
    try:
        hit = _cache.get(year)
        if _fresh(hit, time.monotonic()):
            return hit[2]
        try:
            conn = _connect()
            try:
                if hit and hit[0] == ddb.get_version(conn, year):
                    snap = hit[2]
                else:
                    snap = _snapshot(conn, year, create=False)
                    snap["version"] = snap["config"]["version"]
            finally:
                conn.close()
        except (DrawError, mysql.connector.Error) as e:
            if hit:
                log.warning("draw %s: БД недоступна, отдаём прошлый снимок: %s", year, e)
                _store(year, hit[0], hit[2])  # отметка времени — не долбим БД чаще CACHE_TTL_S
                return hit[2]
            if isinstance(e, DrawError):
                raise
            raise DrawUnavailable("БД недоступна") from e
        _store(year, snap["version"], snap)
    finally:
        lock.release()
    _ensure_probabilities(snap)
    return snap


def _ensure_probabilities(snap: dict) -> None:
    """Пересчёт вероятностей в фоне при смене версии; публика видит последний готовый."""
    year, version = snap["year"], snap["version"]
    with _lock:
        done = _probs.get(year)
        if (done and done[0] == version) or year in _probs_running:
            return
        _probs_running.add(year)

    def run():
        try:
            probs = win_probabilities(snap["applicants"], snap["lists"], snap["config"],
                                      snap["config"]["current_round"], sims=SIMS)
        except Exception:
            log.exception("draw %s: расчёт вероятностей упал (версия %s)", year, version)
            probs = None  # (version, None): не перезапускаем каждую секунду
        with _lock:
            done = _probs.get(year)
            if done is None or done[0] <= version:
                _probs[year] = (version, probs)
            _probs_running.discard(year)

    threading.Thread(target=run, daemon=True).start()


def _last_row(snap: dict) -> dict | None:
    active = [p for p in snap["picks"] if p["cancelled_at"] is None]
    if not active:
        return None
    last_id = max(p["id"] for p in active)
    for c, rows in snap["lists"].items():
        for r in rows:
            if r["pick_id"] == last_id:
                return {**r, "category": c}
    return None


def _aquarium_group(a: dict, rnd: int) -> str:
    if a["format"] == "relay":
        return "relay"
    if rnd >= 4:
        return "individual"
    if a["is_wildcard"]:  # необъявленный WC выглядит как обычный участник своей группы
        return "repeat" if a["starts_individual"] > 0 else "novice"
    return a["pool"]


def public_state(year: int) -> dict:
    """Публично: без признака wild card. Необъявленные WC в аквариуме как обычные участники своей
    группы; вероятности скрыты у всех, пока объявлены не все WC (иначе WC вычислялись бы по цифре).
    Готовый ответ кэшируется по (версия снимка, версия вероятностей) — вызывающий его не меняет."""
    snap = _cached_snapshot(year)
    done = _probs.get(year)
    key = (snap["version"], done[0] if done else None)
    with _lock:
        hit = _state_cache.get(year)
    if hit and hit[0] == key:
        return hit[1]
    state = _build_public_state(year, snap, done)
    with _lock:
        _state_cache[year] = (key, state)
    return state


def _build_public_state(year: int, snap: dict, done) -> dict:
    cfg, rnd = snap["config"], snap["config"]["current_round"]
    probs = done[1] if done else None
    placed = {r["applicant_id"] for rows in snap["lists"].values() for r in rows}
    people = [a for a in snap["applicants"] if a["id"] not in placed]
    if any(a["is_wildcard"] for a in people):
        probs = None
    group_keys = ["relay"] + (["individual"] if rnd >= 4 else ["repeat", "novice"])
    groups = []
    for key in group_keys:
        members = [a for a in people if _aquarium_group(a, rnd) == key]
        label, color = AQUARIUM_GROUPS[key]
        groups.append({"key": key, "label": label, "color": color, "people": len(members),
                       "balls": sum(a["chances"] for a in members)})
    order = {k: i for i, k in enumerate(group_keys)}
    people.sort(key=lambda a: (order[_aquarium_group(a, rnd)], -a["chances"], a["surname"], a["name"]))
    last = _last_row(snap)
    return {
        "year": year, "round": rnd, "round_label": ROUND_LABELS[rnd],
        "capacities": {c: cfg[f"capacity_{c}"] for c in CATEGORIES},
        "lists": {c: [{
            "position": r["position"], "surname": r["surname"], "name": r["name"], "city": r["city"],
            "section": r["section"], "status": PAYMENT_ICONS[r["payment_status"]],
            "is_last": last is not None and r["pick_id"] == last["pick_id"],
        } for r in rows] for c, rows in snap["lists"].items()},
        "aquarium": {"groups": groups, "people": [{
            "surname": a["surname"], "name": a["name"], "city": a["city"], "chances": a["chances"],
            "group": _aquarium_group(a, rnd), "probability": None if probs is None else probs.get(a["id"]),
        } for a in people]},
    }


# ── Фиды vMix ───────────────────────────────────────────────────────────────

def feed_rows(year: int, name: str) -> list[dict]:
    snap = _cached_snapshot(year)
    cfg = snap["config"]
    if name in CATEGORIES:
        cap = cfg[f"capacity_{name}"]
        rows = snap["lists"][name]
        out = [{"num": r["position"], "name": _fio(r), "city": r["city"], "section": SECTION_LABELS[r["section"]]}
               for r in rows]
        for i in range(len(rows) + 1, (cap or 0) + 1):
            out.append({"num": i, "name": "", "city": "", "section": SECTION_LABELS["winner"]})
        if not out:  # места не заданы и список пуст: одна пустая строка для настройки колонок в vMix
            out.append({"num": "", "name": "", "city": "", "section": ""})
        return out
    if name == "wildcard":
        wc = sorted(({**r, "category": c} for c, rows in snap["lists"].items() for r in rows if r["is_wildcard"]),
                    key=lambda r: r["position_key"])
        if not wc:
            return [{"num": "", "name": "", "city": "", "category": ""}]
        return [{"num": i, "name": _fio(r), "city": r["city"], "category": CATEGORY_LABELS[r["category"]]}
                for i, r in enumerate(wc, start=1)]
    if name == "last":
        last = _last_row(snap)
        if last is None:  # одна пустая строка: режиссёр настроит колонки Data Source в vMix до эфира
            return [{"name": "", "city": "", "category": "", "num": "", "round": ""}]
        return [{"name": _fio(last), "city": last["city"], "category": CATEGORY_LABELS[last["category"]],
                 "num": last["position"], "round": ROUND_LABELS[last["round"]]}]
    raise DrawError("нет такого фида")


# Оверлеи для vMix (готовая инфографика по макетам дизайнера): название зачёта на плашке слева.
OVERLAY_LABELS = {"men": "Индивидуальный зачёт. Мужчины", "women": "Индивидуальный зачёт. Женщины",
                  "relay": "Эстафетный зачёт", "wildcard": "Wild card"}


def overlay_rows(year: int, name: str) -> dict:
    """Строки оверлея: все номера сразу (эстафеты — по числу эстафетных заявок, но не меньше мест;
    wild card — по числу wild card), имя — когда вытянут; раздел — победитель / лист ожидания / wild card."""
    if name not in OVERLAY_LABELS:
        raise DrawError("нет такого оверлея")
    snap = _cached_snapshot(year)
    if name == "wildcard":
        drawn = sorted((r for rows in snap["lists"].values() for r in rows if r["is_wildcard"]),
                       key=lambda r: r["position_key"])
        rows = [{"num": i, "name": _fio(r), "section": "wildcard"} for i, r in enumerate(drawn, start=1)]
        total = sum(1 for a in snap["applicants"] if a["is_wildcard"])
        rows += [{"num": i, "name": "", "section": "wildcard"} for i in range(len(rows) + 1, total + 1)]
    else:
        cap = snap["config"][f"capacity_{name}"] or 0
        rows = [{"num": r["position"], "name": _fio(r), "section": r["section"] or "winner"} for r in snap["lists"][name]]
        total = max(cap, sum(1 for a in snap["applicants"] if category_of(a) == name))
        rows += [{"num": i, "name": "", "section": "winner" if i <= cap else "waitlist"}
                 for i in range(len(rows) + 1, total + 1)]
    return {"label": OVERLAY_LABELS[name], "rows": rows}


def to_csv(rows: list[dict], columns: tuple) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(columns), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return "\ufeff" + buf.getvalue()
