"""
Жеребьёвка на локальной БД siberman (.env.local): вытягивание, исправление/отмена,
оплаты с автопередачей, публичные данные без утечки wild card, фиды, маршруты.
Год теста 2090–2099 (от PID), чистится до и после.
"""
import datetime
import io
import os
import threading
import time

import pytest

from app import app
from src.core.auth import api_require_auth
from src.siberman.db import get_siberman_connection
from src.siberman.draw import db as ddb
from src.siberman.draw import service
from src.siberman.draw.service import DrawError, DrawUnavailable

YEAR = 2090 + os.getpid() % 10  # изоляция параллельных прогонов
OTHER_YEAR = 2090 + (YEAR - 2090 + 1) % 10


def _app(surname, fmt="individual", gender="M", pool="novice", wc=False, chances=1, starts=0):
    return {"format": fmt, "surname": surname, "name": "Тест",
            "gender": gender if fmt == "individual" else None, "city": "Красноярск", "pool": pool,
            "is_wildcard": wc, "chances": chances, "starts_individual": starts, "starts_relay": 0,
            "draw_accumulated": 0, "is_local": 0, "duathlon_222": 0, "race_24h": 0, "tattoo": 0,
            "relay_swim": "", "relay_bike": "", "relay_run": "", "wildcard_reason": "", "results_text": ""}


APPS = [
    _app("Вайлдкардов", pool="wildcard", wc=True, starts=2),
    _app("Повторов", pool="repeat", starts=1, chances=3),
    _app("Новиков"), _app("Новикова", gender="F"), _app("Запасной"), _app("Последний"),
    _app("Капитанов", fmt="relay", pool="relay", chances=2),
]


def _clean(conn):
    cur = conn.cursor()
    for t in ("draw_audit", "draw_picks", "draw_applicants", "draw_config"):
        cur.execute(f"DELETE FROM {t} WHERE draw_year=%s", (YEAR,))


@pytest.fixture
def draw():
    conn = get_siberman_connection()
    if conn is None:
        pytest.skip("Нет локальной БД siberman")
    cur = conn.cursor()
    cur.execute("SHOW TABLES LIKE 'draw_picks'")
    if not cur.fetchone():
        conn.close()
        pytest.skip("Миграция 013_draw не применена")
    _clean(conn)
    ddb.replace_applicants(conn, YEAR, APPS)
    ddb.save_config(conn, YEAR, {"capacity_men": 3, "capacity_women": 1, "capacity_relay": 1,
                                 "quota_repeat": 2, "quota_novice": 2,
                                 "auto_invite_until": datetime.date(2098, 1, 1)})
    service.clear_cache()
    yield {a["surname"]: a["id"] for a in ddb.load_applicants(conn, YEAR)}
    _clean(conn)
    conn.close()
    service.clear_cache()


def _men_after_wildcard_and_three(ids):
    service.pick(YEAR, ids["Вайлдкардов"])
    service.set_round(YEAR, 4)
    for s in ("Новиков", "Запасной", "Последний"):
        service.pick(YEAR, ids[s])


def test_wildcard_round_and_refusals(draw):
    r = service.pick(YEAR, draw["Вайлдкардов"])
    assert (r["category"], r["position"]) == ("men", 1)
    with pytest.raises(DrawError, match="уже вытянут — Мужчины №1"):
        service.pick(YEAR, draw["Вайлдкардов"])
    with pytest.raises(DrawError, match="не в этом раунде"):
        service.pick(YEAR, draw["Новиков"])


def test_round2_repeat_quota_defers_to_round4(draw):
    service.save_config(YEAR, {"quota_repeat": 1})
    service.pick(YEAR, draw["Вайлдкардов"])
    service.set_round(YEAR, 2)
    with pytest.raises(DrawError, match=r"квота повторных набрана \(1/1\)"):
        service.pick(YEAR, draw["Повторов"])
    blocks = {a["surname"]: a["block"] for a in service.admin_state(YEAR)["applicants"]}
    assert "квота повторных набрана" in blocks["Повторов"] and blocks["Новиков"].startswith("не в этом раунде")
    service.set_round(YEAR, 4)
    assert service.pick(YEAR, draw["Повторов"])["position"] == 2


def test_round3_full_gender_goes_to_waitlist_without_quota(draw):
    service.save_config(YEAR, {"capacity_women": 0, "quota_novice": 1})
    service.set_round(YEAR, 3)
    service.pick(YEAR, draw["Новикова"])
    women = service.admin_state(YEAR)["lists"]["women"]
    assert [(r["surname"], r["section"]) for r in women] == [("Новикова", "waitlist")]
    service.pick(YEAR, draw["Новиков"])                       # женщина из листа ожидания квоту не заняла
    with pytest.raises(DrawError, match=r"квота новичков набрана \(1/1\)"):
        service.pick(YEAR, draw["Последний"])


def test_replace_keeps_position_cancel_shifts(draw):
    service.set_round(YEAR, 4)
    a = service.pick(YEAR, draw["Новиков"])
    b = service.pick(YEAR, draw["Запасной"])
    service.pick(YEAR, draw["Последний"])
    service.replace(YEAR, a["pick_id"], draw["Повторов"])
    assert [r["surname"] for r in service.admin_state(YEAR)["lists"]["men"]] == ["Повторов", "Запасной", "Последний"]
    service.cancel(YEAR, b["pick_id"])
    men = service.admin_state(YEAR)["lists"]["men"]
    assert [(r["position"], r["surname"]) for r in men] == [(1, "Повторов"), (2, "Последний")]
    assert service.pick(YEAR, draw["Запасной"])["position"] == 3


def test_replace_other_category_keeps_draw_order(draw):
    # вытянули мужчину, а ввели женщину: исправленный встаёт к мужчинам на свою очередь, соседи сдвигаются
    service.save_config(YEAR, {"capacity_men": 2})
    service.set_round(YEAR, 4)
    service.pick(YEAR, draw["Новиков"])
    wrong = service.pick(YEAR, draw["Новикова"])
    service.pick(YEAR, draw["Последний"])
    r = service.replace(YEAR, wrong["pick_id"], draw["Запасной"])
    assert (r["category"], r["position"]) == ("men", 2)
    assert r["shifted"] == [{"surname": "Последний", "name": "Тест", "category": "men", "position": 3,
                             "section": "waitlist"}]
    lists = service.admin_state(YEAR)["lists"]
    assert [r["surname"] for r in lists["men"]] == ["Новиков", "Запасной", "Последний"] and lists["women"] == []
    back = service.replace(YEAR, r["pick_id"], draw["Новикова"])      # и обратно
    assert (back["category"], back["position"], back["shifted"][0]["section"]) == ("women", 1, "winner")


def test_replace_still_checks_round_of_the_record(draw):
    service.set_round(YEAR, 1)
    team = service.pick(YEAR, draw["Капитанов"])
    with pytest.raises(DrawError, match="не в этом раунде"):
        service.replace(YEAR, team["pick_id"], draw["Новиков"])


def test_backup_excel_has_whole_course(draw):
    from openpyxl import load_workbook
    service.set_round(YEAR, 4)
    wrong = service.pick(YEAR, draw["Новикова"])
    service.pick(YEAR, draw["Последний"])
    service.replace(YEAR, wrong["pick_id"], draw["Новиков"])
    wb = load_workbook(io.BytesIO(service.draw_backup(YEAR)))
    assert wb.sheetnames == ["Ход жеребьёвки", "Итог", "Журнал действий", "Параметры"]
    course = [[c.value for c in row] for row in wb["Ход жеребьёвки"].iter_rows(min_row=2)]
    assert [(r[3], r[5], r[6], r[9]) for r in course] == [
        ("Новикова Тест", "Женщины", None, "исправлено → Новиков Тест"),
        ("Новиков Тест", "Мужчины", 1, "в списке"),           # исправление — на очереди исправленного
        ("Последний Тест", "Мужчины", 2, "в списке")]
    final = [[c.value for c in row] for row in wb["Итог"].iter_rows(min_row=2)]
    assert [(r[0], r[1], r[3]) for r in final] == [("Мужчины", 1, "Новиков Тест"), ("Мужчины", 2, "Последний Тест")]
    log = [[c.value for c in row] for row in wb["Журнал действий"].iter_rows(min_row=2)]
    assert [(r[1], r[2], r[3], r[4]) for r in log] == [("исправление", "Новиков Тест", "Новикова Тест", "Новиков Тест")]


def test_unpaid_auto_invites_next_and_suggests_revert(draw):
    _men_after_wildcard_and_three(draw)          # мест 3 → «Последний» №4 в листе ожидания
    r = service.set_payment(YEAR, draw["Новиков"], "unpaid", today=datetime.date(2097, 11, 1))
    assert (r["auto_invited"]["surname"], r["auto_invited"]["position"]) == ("Последний", 4)
    status = {a["surname"]: a["payment_status"] for a in service.admin_state(YEAR)["applicants"]}
    assert status["Последний"] == "pending"
    back = service.set_payment(YEAR, draw["Новиков"], "paid", today=datetime.date(2097, 11, 2))
    assert back["suggest_revert"]["surname"] == "Последний"


def test_auto_invite_off_after_deadline(draw):
    _men_after_wildcard_and_three(draw)
    r = service.set_payment(YEAR, draw["Новиков"], "unpaid", today=datetime.date(2098, 1, 1))
    assert r["auto_invited"] is None


def test_waitlist_empty_warning(draw):
    service.set_round(YEAR, 4)
    for s in ("Новиков", "Запасной", "Последний"):
        service.pick(YEAR, draw[s])
    r = service.set_payment(YEAR, draw["Новиков"], "unpaid", today=datetime.date(2097, 11, 1))
    assert r["waitlist_empty"] is True and r["auto_invited"] is None


def test_payment_requires_drawn(draw):
    with pytest.raises(DrawError, match="ещё не вытянут"):
        service.set_payment(YEAR, draw["Новиков"], "paid")


def _wait_probs():
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline:
        done = service._probs.get(YEAR)
        if done and done[1] is not None:
            return
        time.sleep(0.05)


def test_public_aquarium_shows_unrevealed_wildcard_as_ordinary(draw):
    service.public_state(YEAR)
    _wait_probs()
    pub = service.public_state(YEAR)
    people = {p["surname"]: p for p in pub["aquarium"]["people"]}
    assert people["Вайлдкардов"]["group"] == "repeat"
    assert not {"is_wildcard", "pool"} & set(people["Вайлдкардов"])
    assert next(g for g in pub["aquarium"]["groups"] if g["key"] == "repeat")["people"] == 2
    assert all(p["probability"] is None for p in pub["aquarium"]["people"])
    # до первого вытягивания фид — одна пустая строка с ключами (колонки Data Source в vMix)
    assert service.feed_rows(YEAR, "wildcard") == [{"num": "", "name": "", "city": "", "category": ""}]
    service.pick(YEAR, draw["Вайлдкардов"])
    service.clear_cache()
    service.public_state(YEAR)
    _wait_probs()
    pub = service.public_state(YEAR)
    assert "Вайлдкардов" not in [p["surname"] for p in pub["aquarium"]["people"]]
    assert pub["lists"]["men"][0]["surname"] == "Вайлдкардов"
    assert "is_wildcard" not in pub["lists"]["men"][0]
    assert any(p["probability"] is not None for p in pub["aquarium"]["people"])
    assert [r["name"] for r in service.feed_rows(YEAR, "wildcard")] == ["Вайлдкардов Тест"]


def test_feed_padded_to_capacity_and_csv(draw):
    rows = service.feed_rows(YEAR, "men")
    assert [r["num"] for r in rows] == [1, 2, 3] and all(r["name"] == "" for r in rows)
    assert service.to_csv(rows, service.FEED_COLUMNS["men"]).startswith("\ufeffnum,name,city,section\n")


def test_import_blocked_after_start(draw, monkeypatch):
    monkeypatch.setattr(service, "parse_applicants", lambda b: (APPS, []))
    service.pick(YEAR, draw["Вайлдкардов"])
    with pytest.raises(DrawError, match="уже началась"):
        service.import_applicants(YEAR, b"x", dry_run=False)
    checked = service.import_applicants(YEAR, b"x", dry_run=True)
    assert checked["written"] is False and checked["blocked"] is True


def test_import_dry_run_shows_diff_against_loaded(draw, monkeypatch):
    new = [{**a, "chances": 4, "draw_accumulated": 1} if a["surname"] == "Повторов" else a
           for a in APPS if a["surname"] != "Последний"] + [_app("Добавленный")]
    monkeypatch.setattr(service, "parse_applicants", lambda b: (new, []))
    d = service.import_applicants(YEAR, b"x", dry_run=True)
    assert (d["written"], d["blocked"]) == (False, False)
    assert [x["fio"] for x in d["diff"]["added"]] == ["Добавленный Тест"]
    assert [x["fio"] for x in d["diff"]["removed"]] == ["Последний Тест"]
    assert [x["changes"] for x in d["diff"]["changed"]] == [[{"field": "chances", "old": 3, "new": 4},
                                                             {"field": "draw_accumulated", "old": 0, "new": 1}]]
    stored = {a["surname"]: a["chances"] for a in service.admin_state(YEAR)["applicants"]}
    assert stored["Повторов"] == 3 and "Добавленный" not in stored          # проверка ничего не записала


def test_reset_needs_year(draw):
    service.pick(YEAR, draw["Вайлдкардов"])
    with pytest.raises(DrawError):
        service.reset(YEAR, "2028")
    service.reset(YEAR, str(YEAR))
    assert service.admin_state(YEAR)["lists"]["men"] == []


def test_summary_has_history_and_placement(draw):
    service.pick(YEAR, draw["Вайлдкардов"])
    s = service.applicant_summary(YEAR, draw["Вайлдкардов"])
    assert s["placement"]["category"] == "men" and s["placement"]["position"] == 1
    assert set(s["history"]) == {"individual", "relay", "finished"}


def test_routes(draw, client):
    app.dependency_overrides[api_require_auth] = lambda: "testuser"
    try:
        r = client.post(f"/api/siberman/admin/draw/{YEAR}/pick", json={"applicant_id": draw["Вайлдкардов"]})
        assert r.status_code == 200 and r.json()["position"] == 1
        r = client.post(f"/api/siberman/admin/draw/{YEAR}/pick", json={"applicant_id": draw["Вайлдкардов"]})
        assert r.status_code == 409 and "уже вытянут" in r.json()["detail"]
    finally:
        app.dependency_overrides.pop(api_require_auth, None)
    service.clear_cache()
    assert client.get(f"/api/siberman/draw/{YEAR}/state").json()["lists"]["men"][0]["surname"] == "Вайлдкардов"
    csv_r = client.get(f"/api/siberman/draw/{YEAR}/feed/men.csv")
    assert csv_r.status_code == 200 and csv_r.content.startswith(b"\xef\xbb\xbf")
    assert client.get(f"/api/siberman/draw/{YEAR}/feed/nope.json").status_code == 404
    assert client.get(f"/draw/{YEAR}/results", headers={"host": "siberman.localhost"}).status_code == 200
    assert client.get(f"/draw/{YEAR}/results").status_code == 404
    assert client.get(f"/api/siberman/admin/draw/{YEAR}/state").status_code == 401


# ── Ревью: настройки, кэш, сбой БД, оплаты, права ───────────────────────────

def _status(ids, surname):
    return next(a for a in service.admin_state(YEAR)["applicants"] if a["id"] == ids[surname])


def test_last_feed_empty_row_before_first_pick(draw):
    assert service.feed_rows(YEAR, "last") == [{"name": "", "city": "", "category": "", "num": "", "round": ""}]


def test_partial_config_post_keeps_other_fields(draw, client):
    app.dependency_overrides[api_require_auth] = lambda: "testuser"
    try:
        r = client.post(f"/api/siberman/admin/draw/{YEAR}/config", json={"capacity_men": 5})
        assert r.status_code == 200
        cfg = service.admin_state(YEAR)["config"]
        assert cfg["capacity_men"] == 5 and cfg["capacity_women"] == 1 and cfg["quota_repeat"] == 2
        assert cfg["auto_invite_until"] == datetime.date(2098, 1, 1)
        assert client.post(f"/api/siberman/admin/draw/{YEAR}/config", json={"capacity_men": 1001}).status_code == 409
    finally:
        app.dependency_overrides.pop(api_require_auth, None)


def test_pick_concurrent_same_applicant(draw):
    results = []

    def go():
        try:
            results.append(service.pick(YEAR, draw["Вайлдкардов"]))
        except DrawError as e:
            results.append(e)
    threads = [threading.Thread(target=go) for _ in range(4)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    errors = [r for r in results if isinstance(r, DrawError)]
    assert len(results) == 4 and len(errors) == 3 and all("уже вытянут" in str(e) for e in errors)


def test_pick_concurrent_different_applicants_no_duplicate_numbers(draw):
    service.set_round(YEAR, 4)
    names = ("Новиков", "Запасной", "Последний", "Повторов")
    errors = []

    def go(n):
        try:
            service.pick(YEAR, draw[n])
        except Exception as e:  # noqa: BLE001
            errors.append(e)
    threads = [threading.Thread(target=go, args=(n,)) for n in names]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert not errors
    assert sorted(r["position"] for r in service.admin_state(YEAR)["lists"]["men"]) == [1, 2, 3, 4]


def test_public_cache_invalidated_by_version_without_clear(draw, monkeypatch):
    monkeypatch.setattr(service, "CACHE_TTL_S", 0)
    assert service.public_state(YEAR)["lists"]["men"] == []
    service.pick(YEAR, draw["Вайлдкардов"])
    assert service.public_state(YEAR)["lists"]["men"][0]["surname"] == "Вайлдкардов"


def test_db_failure_serves_stale_snapshot_or_unavailable(draw, monkeypatch):
    monkeypatch.setattr(service, "CACHE_TTL_S", 0)
    real = service.get_siberman_connection
    monkeypatch.setattr(service, "get_siberman_connection", lambda: None)
    with pytest.raises(DrawUnavailable):
        service.public_state(YEAR)               # кэш пуст (clear_cache в фикстуре)
    monkeypatch.setattr(service, "get_siberman_connection", real)
    warm = service.public_state(YEAR)
    monkeypatch.setattr(service, "get_siberman_connection", lambda: None)
    stale = service.public_state(YEAR)           # БД лежит — отдаём старый снимок
    # (вероятности аквариума досчитываются в фоне, поэтому сравниваем списки и раунд)
    assert (stale["lists"], stale["round"]) == (warm["lists"], warm["round"])


def test_db_down_route_returns_503(draw, client, monkeypatch):
    monkeypatch.setattr(service, "get_siberman_connection", lambda: None)
    assert client.get(f"/api/siberman/draw/{YEAR}/state").status_code == 503
    assert client.get(f"/api/siberman/draw/{YEAR}/feed/men.json").status_code == 503


def test_public_read_does_not_create_config_row(draw):
    conn = get_siberman_connection()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM draw_config WHERE draw_year=%s", (OTHER_YEAR,))
        pub = service.public_state(OTHER_YEAR)
        assert pub["round"] == 0 and all(v is None for v in pub["capacities"].values())
        assert all(rows == [] for rows in pub["lists"].values()) and pub["aquarium"]["people"] == []
        assert service.feed_rows(OTHER_YEAR, "last")[0]["name"] == ""
        cur.execute("SELECT COUNT(*) FROM draw_config WHERE draw_year=%s", (OTHER_YEAR,))
        assert cur.fetchone()[0] == 0
    finally:
        conn.close()
        service.clear_cache()


def test_empty_category_feed_without_capacity_is_one_stub_row(draw):
    conn = get_siberman_connection()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM draw_config WHERE draw_year=%s", (OTHER_YEAR,))
        for name in ("men", "women", "relay"):
            assert service.feed_rows(OTHER_YEAR, name) == [
                {"num": "", "name": "", "city": "", "section": ""}]
    finally:
        conn.close()
        service.clear_cache()


def test_public_state_cached_while_versions_unchanged(draw):
    service.public_state(YEAR)
    _wait_probs()
    first = service.public_state(YEAR)
    assert service.public_state(YEAR) is first
    service.pick(YEAR, draw["Вайлдкардов"])
    service.clear_cache()
    assert service.public_state(YEAR) is not first


def test_failed_probabilities_not_restarted(draw, monkeypatch):
    calls = []

    def boom(*a, **k):
        calls.append(1)
        raise RuntimeError("sim")
    monkeypatch.setattr(service, "win_probabilities", boom)
    monkeypatch.setattr(service, "CACHE_TTL_S", 0)
    service.public_state(YEAR)
    for _ in range(100):
        if service._probs.get(YEAR):
            break
        time.sleep(0.05)
    assert service._probs[YEAR][1] is None
    service.public_state(YEAR)
    time.sleep(0.2)
    assert len(calls) == 1


def test_auto_invite_keeps_comment_and_audits_comment_change(draw):
    _men_after_wildcard_and_three(draw)
    today = datetime.date(2097, 11, 1)
    service.set_payment(YEAR, draw["Последний"], "none", "звонила мама", today=today)
    service.set_payment(YEAR, draw["Новиков"], "unpaid", today=today)
    # «Последний» в листе ожидания со своим комментарием, статус none → его и зовут, комментарий цел
    assert _status(draw, "Последний")["payment_status"] == "pending"
    assert _status(draw, "Последний")["payment_comment"] == "звонила мама"
    service.set_payment(YEAR, draw["Новиков"], "unpaid", "договорились", today=today)
    actions = [x["action"] for x in service.applicant_summary(YEAR, draw["Новиков"])["audit"]]
    assert actions == ["payment", "payment_comment"]


def test_unpaid_paid_unpaid_does_not_invite_second(draw):
    _men_after_wildcard_and_three(draw)
    today = datetime.date(2097, 11, 1)
    first = service.set_payment(YEAR, draw["Новиков"], "unpaid", today=today)
    assert first["auto_invited"]["surname"] == "Последний"
    service.set_payment(YEAR, draw["Новиков"], "paid", today=today)
    again = service.set_payment(YEAR, draw["Новиков"], "unpaid", today=today)
    assert again["auto_invited"] is None and again["already_invited"]["surname"] == "Последний"


def test_cancel_and_replace_reset_payment_status(draw):
    service.set_round(YEAR, 4)
    a = service.pick(YEAR, draw["Новиков"])
    b = service.pick(YEAR, draw["Запасной"])
    service.set_payment(YEAR, draw["Новиков"], "paid")
    service.set_payment(YEAR, draw["Запасной"], "paid")
    service.cancel(YEAR, a["pick_id"])
    assert _status(draw, "Новиков")["payment_status"] == "none"
    service.replace(YEAR, b["pick_id"], draw["Повторов"])
    assert _status(draw, "Запасной")["payment_status"] == "none"
    assert "payment" in [x["action"] for x in service.applicant_summary(YEAR, draw["Запасной"])["audit"]]


ADMIN_CALLS = [
    ("get", "state", None), ("get", "applicant/1", None), ("post", "round", {"round": 1}),
    ("post", "pick", {"applicant_id": 1}), ("post", "picks/1/cancel", None),
    ("post", "picks/1/replace", {"applicant_id": 1}), ("get", "backup.xlsx", None),
    ("post", "payment", {"applicant_id": 1, "status": "paid"}), ("post", "config", {}),
    ("post", "reset", {"confirm": "x"}),
]


def _call(client, method, year, path, body):
    kw = {"json": body} if body is not None else {}
    return getattr(client, method)(f"/api/siberman/admin/draw/{year}/{path}", **kw)


@pytest.mark.parametrize("method,path,body", ADMIN_CALLS)
def test_admin_endpoints_require_auth(client, method, path, body):
    assert _call(client, method, YEAR, path, body).status_code == 401


@pytest.mark.parametrize("bad_year", [0, 99999])
@pytest.mark.parametrize("method,path,body", ADMIN_CALLS)
def test_admin_endpoints_reject_bad_year(client, method, path, body, bad_year):
    app.dependency_overrides[api_require_auth] = lambda: "testuser"
    try:
        assert _call(client, method, bad_year, path, body).status_code == 404
    finally:
        app.dependency_overrides.pop(api_require_auth, None)


def test_import_too_large_rejected(client):
    app.dependency_overrides[api_require_auth] = lambda: "testuser"
    try:
        big = b"x" * (5 * 1024 * 1024 + 1)
        r = client.post(f"/api/siberman/admin/draw/{YEAR}/import", files={"file": ("a.xlsx", big)})
        assert r.status_code == 409 and "слишком большой" in r.json()["detail"]
    finally:
        app.dependency_overrides.pop(api_require_auth, None)


def test_login_returns_to_requested_page(client):
    """Вход с /admin/draw возвращает на /admin/draw, а не на /admin."""
    from urllib.parse import parse_qs, urlparse
    from src.config import settings
    host = {"host": "siberman.localhost"}
    r = client.get("/admin/draw", headers=host, follow_redirects=False)
    assert r.status_code == 302
    loc = urlparse(r.headers["location"])
    assert loc.path == "/login" and parse_qs(loc.query)["next"] == ["/admin/draw"]
    r = client.post("/login", headers=host, follow_redirects=False,
                    data={"username": settings.ADMIN_USERNAME, "password": settings.ADMIN_PASSWORD, "next": "/admin/draw"})
    assert r.status_code == 302 and r.headers["location"] == "/admin/draw"
    r = client.post("/login", headers=host, follow_redirects=False,
                    data={"username": settings.ADMIN_USERNAME, "password": settings.ADMIN_PASSWORD, "next": "//evil.example"})
    assert r.headers["location"] == "/admin"
    client.cookies.clear()

def test_host_page_and_state(draw, client):
    host, token, bad = {"host": "siberman.localhost"}, service.host_token(YEAR), "0" * 32
    assert client.get(f"/draw/{YEAR}/host/{bad}", headers=host).status_code == 404
    r = client.get(f"/draw/{YEAR}/host/{token}", headers=host)
    assert r.status_code == 200 and "noindex" in r.headers.get("x-robots-tag", "")
    assert client.get(f"/draw/{YEAR}/host/{token}").status_code == 404                 # чужой домен
    assert client.get(f"/api/siberman/draw/{YEAR}/host/{bad}/state").status_code == 404
    service.pick(YEAR, draw["Вайлдкардов"])
    service.set_round(YEAR, 4)
    service.pick(YEAR, draw["Новиков"])
    service.clear_cache()
    st = client.get(f"/api/siberman/draw/{YEAR}/host/{token}/state").json()
    assert st["last"] == draw["Новиков"]
    assert [r["applicant_id"] for r in st["lists"]["men"]] == [draw["Вайлдкардов"], draw["Новиков"]]
    wc = st["people"][str(draw["Вайлдкардов"])]
    assert wc["is_wildcard"] and wc["chances"] == 1 and "history" in wc and "payment_status" not in wc
    service.cancel(YEAR, service.admin_state(YEAR)["journal"][0]["pick_id"])
    service.clear_cache()
    assert client.get(f"/api/siberman/draw/{YEAR}/host/{token}/state").json()["last"] == draw["Вайлдкардов"]
    assert service.admin_state(YEAR)["host_link"] == f"/draw/{YEAR}/host/{token}"


def _tilda_csv(rows):
    head = "surname;Name;Email;Phone;Date;format"
    return ("\ufeff" + "\n".join([head, *(";".join(r) for r in rows)])).encode("utf-8")


def test_contacts_and_payment_export(draw, client):
    import io
    import openpyxl
    data = _tilda_csv([
        ("Вайлдкардов", "Тест", "old@ex.ru", "+7 900 000-00-00", "2026-09-01 10:00:00", "Индивидуальное участие"),
        ("Вайлдкардов", "Тест", "wc@ex.ru", "+7 900 111-11-11", "2026-09-02 10:00:00", "Индивидуальное участие"),
        ("Последний", "Тест", "last@ex.ru", "+7 900 222-22-22", "2026-09-01 10:00:00", "Индивидуальное участие"),
        ("Капитанов", "Тест", "cap@ex.ru", "+7 900 333-33-33", "2026-09-01 10:00:00", "Эстафетная команда"),
    ])
    rep = service.import_contacts(YEAR, data, "test")
    assert (rep["people"], rep["duplicates"]) == (3, 1)
    conn = get_siberman_connection()
    try:                                                      # новая версия таблицы шансов — контакты остаются
        ddb.replace_applicants(conn, YEAR, APPS)
        ids = {a["surname"]: a["id"] for a in ddb.load_applicants(conn, YEAR)}
    finally:
        conn.close()
    service.clear_cache()
    service.pick(YEAR, ids["Вайлдкардов"])
    service.set_round(YEAR, 4)
    for s in ("Новиков", "Запасной", "Последний"):            # мест мужчин 3: Последний — лист ожидания
        service.pick(YEAR, ids[s])
    assert "Последний Тест" not in [r["fio"] for r in service.payment_rows(YEAR)]
    service.set_payment(YEAR, ids["Последний"], "pending")    # приглашён из листа ожидания
    rows = {r["fio"]: r for r in service.payment_rows(YEAR)}
    assert set(rows) == {"Вайлдкардов Тест", "Новиков Тест", "Запасной Тест", "Последний Тест"}
    assert (rows["Вайлдкардов Тест"]["email"], rows["Вайлдкардов Тест"]["wildcard"]) == ("wc@ex.ru", "да")
    assert rows["Новиков Тест"]["email"] == "" and rows["Последний Тест"]["section"] == "Лист ожидания"
    assert "contacts" not in str(service.public_state(YEAR)) and "wc@ex.ru" not in str(service.public_state(YEAR))
    assert "wc@ex.ru" not in str(service.host_state(YEAR))

    assert client.get(f"/api/siberman/admin/draw/{YEAR}/export.xlsx").status_code == 401
    app.dependency_overrides[api_require_auth] = lambda: "testuser"
    try:
        r = client.get(f"/api/siberman/admin/draw/{YEAR}/export.xlsx")
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        table = [[c.value for c in row] for row in ws.iter_rows()]
        assert table[0][:7] == ["Зачёт", "№", "Раздел", "Фамилия Имя", "Город", "E-mail", "Телефон"]
        assert ["wc@ex.ru"] == [row[5] for row in table[1:] if row[3] == "Вайлдкардов Тест"]
        assert client.post(f"/api/siberman/admin/draw/{YEAR}/contacts",
                           files={"file": ("z.csv", b"bad", "text/csv")}).status_code == 409
    finally:
        app.dependency_overrides.pop(api_require_auth, None)

def test_overlay_pages_and_rows(draw, client):
    host = {"host": "siberman.localhost"}
    for name in ("men", "women", "relay", "wildcard"):
        assert client.get(f"/draw/{YEAR}/overlay/{name}", headers=host).status_code == 200
    assert client.get(f"/draw/{YEAR}/overlay/nope", headers=host).status_code == 404
    assert client.get(f"/draw/{YEAR}/overlay/relay").status_code == 404                  # чужой домен
    relay = service.overlay_rows(YEAR, "relay")
    assert relay["label"] == "Эстафетный зачёт"
    assert relay["rows"] == [{"num": 1, "name": "", "section": "winner"}]     # 1 эстафета, 1 место — номер есть, имени нет
    assert service.overlay_rows(YEAR, "wildcard")["rows"] == [{"num": 1, "name": "", "section": "wildcard"}]
    service.pick(YEAR, draw["Вайлдкардов"])
    service.clear_cache()
    assert service.overlay_rows(YEAR, "wildcard")["rows"] == [{"num": 1, "name": "Вайлдкардов Тест", "section": "wildcard"}]
    service.set_round(YEAR, 1)
    service.pick(YEAR, draw["Капитанов"])
    service.clear_cache()
    data = client.get(f"/api/siberman/draw/{YEAR}/overlay/relay").json()
    assert data["rows"] == [{"num": 1, "name": "Капитанов Тест", "section": "winner"}]
    men = service.overlay_rows(YEAR, "men")
    assert men["label"] == "Индивидуальный зачёт. Мужчины" and [r["section"] for r in men["rows"]][:3] == ["winner"] * 3
    assert client.get(f"/api/siberman/draw/{YEAR}/overlay/nope").status_code == 404
