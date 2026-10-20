"""Статистика страницы сертификата (src/analytics/medals_stats.py): лог nginx → сводка и Excel."""
import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from src.analytics import medals_stats as ms

KRSK = timezone(timedelta(hours=7))


def write(path, lines):
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def test_read_keeps_only_known_events_and_frames(tmp_path):
    log = tmp_path / "stat.log"
    write(log, ["2026-10-22T10:00:00+07:00\topen\t\t1.1.1.1",
                "2026-10-22T10:01:00+07:00\tdownload\tmedal-russia\t1.1.1.1",
                "2026-10-22T10:02:00+07:00\tdownload\t<script>\t2.2.2.2",       # чужая рамка → «не указана»
                "2026-10-22T10:03:00+07:00\thack\tmedal-russia\t3.3.3.3",        # чужое событие — мимо
                "мусор без табуляций",
                "not-a-date\topen\t\t4.4.4.4"])
    rows = ms.read(log)
    assert [(e, f) for _, e, f, _ in rows] == [("open", ""), ("download", "medal-russia"), ("download", "")]


def test_summary_counts_frames_hours_visitors_conversion(tmp_path):
    log = tmp_path / "stat.log"
    write(log, ["2026-10-22T10:00:00+07:00\topen\t\t1.1.1.1",
                "2026-10-22T10:05:00+07:00\topen\t\t1.1.1.1",
                "2026-10-22T10:10:00+07:00\topen\t\t2.2.2.2",
                "2026-10-22T10:20:00+07:00\tdownload\tmedal-krasnoyarsk\t1.1.1.1",
                "2026-10-22T11:00:00+07:00\tshare\tmedal-russia\t2.2.2.2"])
    s = ms.summary(ms.read(log))
    assert s["totals"] == {"open": 3, "download": 1, "share": 1}
    assert s["visitors"] == 2 and s["conversion"] == 50.0       # из двух посетителей скачал один
    frames = {r["frame"]: (r["download"], r["share"]) for r in s["by_frame"]}
    assert frames == {"medal-krasnoyarsk": (1, 0), "medal-russia": (0, 1)}
    assert [(r["hour"], r["open"], r["share"]) for r in s["by_hour"]] == [("2026-10-22 11:00", 0, 1), ("2026-10-22 10:00", 3, 0)]


def test_empty_log_and_append_roundtrip(tmp_path):
    log = tmp_path / "logs" / "stat.log"
    assert ms.summary(ms.read(log))["conversion"] is None
    ms.append("download", "medal-russia", "5.5.5.5", path=log, now=datetime(2026, 10, 22, 9, 0, tzinfo=KRSK))
    ms.append("open\tевил", "", "6.6.6.6", path=log)              # табуляция в событии не ломает строку
    rows = ms.read(log)
    assert [(ts.hour, e, f, ip) for ts, e, f, ip in rows] == [(9, "download", "medal-russia", "5.5.5.5")]


def test_xlsx_has_sheets_and_no_ip(tmp_path):
    log = tmp_path / "stat.log"
    write(log, ["2026-10-22T10:00:00+07:00\topen\t\t10.9.8.7",
                "2026-10-22T10:20:00+07:00\tdownload\tmedal-russia\t10.9.8.7"])
    wb = load_workbook(io.BytesIO(ms.to_xlsx(ms.read(log))))
    assert wb.sheetnames == ["Итоги", "По рамкам", "По часам", "События"]
    values = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row]
    assert "10.9.8.7" not in " ".join(values)
    assert [c.value for c in wb["События"][3]] == ["2026-10-22 10:20:00", "Скачали сертификат", "Оборот"]
