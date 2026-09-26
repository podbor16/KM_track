import pytest

from scripts.clean_clients import Names, canonical_fio, find_groups, resolve_group, title, tokens, translit

_ROWS = (
    [("Иванов", n, "Мужчина") for n in ("Иван", "Андрей", "Александр", "Дмитрий", "Евгений", "Сергей", "Владимир",
                                        "Константин", "Максим", "Петр", "Денис", "Олег", "Лука", "Егор")] * 6
    + [("Иванова", n, "Женщина") for n in ("Мария", "Елена", "Наталия", "Ольга", "Екатерина", "Ксения", "Алина", "Ольга", "Самсон", "Игорь", "Есения", "Таисия",
                                           "Анна", "Татьяна", "Юлия", "Светлана", "Дарья")] * 6
    + [(s, "Иван", "Мужчина") for s in ("Молодцов", "Молодцов", "Буров", "Буров", "Кольга", "Кольга")]
)
NAMES = Names(_ROWS)


@pytest.mark.parametrize("lat, cyr", [
    ("Aleksandr", "александр"), ("Dmitry", "дмитрий"), ("Iuliia", "юлия"), ("Oshchepkov", "ощепков"),
    ("Zhukov", "жуков"), ("Andrey", "андрей"), ("Vyacheslav", "вячеслав"),
])
def test_translit(lat, cyr):
    assert translit(lat) == cyr


def test_tokens_cleanup():
    assert tokens("⁠Петрова") == ["Петрова"]
    assert tokens("Деменкова ✅") == ["Деменкова"]
    assert tokens("Петров- Дельверс") == ["Петров-Дельверс"]
    assert tokens("Нинa") == ["Нина"]            # латинская a
    assert tokens("ЛаZученко") == ["Лазученко"]
    assert tokens("Пи́санов") == ["Писанов"]
    assert tokens("Ксения69_") == ["Ксения"]


def test_title():
    assert title("КРАВЧУК") == "Кравчук"
    assert title("мария-луиза") == "Мария-Луиза"
    assert title("МакКой") == "МакКой"


@pytest.mark.parametrize("s, n, sex, exp", [
    ("Артур_нет", "Иванов", "", None),
    ("Андрей", "Буров", "М", ("Буров", "Андрей")),                               # перестановка
    ("Буров", "Андрей Буров", "М", ("Буров", "Андрей")),                         # фамилия в имени
    ("Буров Андрей Сергеевич", "Буров Андрей Сергеевич", "М", ("Буров", "Андрей")),
    ("Иванова", "Мария Сергеевна", "Ж", ("Иванова", "Мария")),                   # отчество
    ("Наталия Валисевич", "Наталия Валисевич", "Ж", ("Валисевич", "Наталия")),   # -вич — фамилия
    ("Иванов Андрей", "Лука", "М", ("Иванов", "Лука")),                          # имя родителя в фамилии
    ("Dmitrii Burmakin", "Лука", "М", ("Бурмакин", "Лука")),
    ("КРАВЧУК", "мария", "Ж", ("Кравчук", "Мария")),
    ("Molodsov", "Adnrei", "М", ("Молодсов", "Андрей")),                        # транслит, перестановка букв в имени
    ("Ivanova", "Mariya", "Ж", ("Иванова", "Мария")),                            # фамилия — только транслит
    ("Troyakova", "Elena", "Ж", ("Троякова", "Елена")),
    ("Rybkin", "Igor", "М", ("Рыбкин", "Игорь")),
    ("Ermakova", "Olga", "Ж", ("Ермакова", "Ольга")),
    ("Sander", "Darya", "Ж", ("Сандер", "Дарья")),
    ("Slawson", "Thomas", "М", ("Slawson", "Thomas")),                          # иностранное — не «Самсон»
    ("Loginov Egor Sergeevich", "Loginov Egor Sergeevich", "М", ("Логинов", "Егор")),
    ("Моrotskiy", "Denis", "М", ("Мороцкий", "Денис")),
    ("Dmitry Khlebnikov", "Dmitry Khlebnikov", "Ж", ("Хлебников", "Дмитрий")),   # точное имя — пол не важен
    ("Богданкевич", "Ратибор", "М", ("Богданкевич", "Ратибор")),               # -вич без имени перед — фамилия
    ("Таисия", "Юрьевна", "Ж", ("", "Таисия")),                                 # фамилии нет
    ("Омарова", "Алмазная Крошка", "Ж", ("", "")),                              # не разобрать
    ("Ястребцова", "Есения - Принцесса", "Ж", ("Ястребцова", "Есения")),
    ("Баранников", "ЕП", "М", ("Баранников", "")),
    ("Konstantin", "Kudashkin", "М", ("Кудашкин", "Константин")),               # транслит + перестановка
    ("Retief", "Nicci", "Ж", ("Retief", "Nicci")),                               # иностранное ФИ
    ("Pelser", "Petra", "Ж", ("Pelser", "Petra")),                               # «Петр» — другой пол
    ("Прусаков Владимир Сергеевич Prusakov Vladimir", "Прусаков Владимир Сергеевич Prusakov Vladimir", "М",
     ("Прусаков", "Владимир")),
    ("Крысанов", "Алексей К.", "М", ("Крысанов", "Алексей")),                     # инициал
    ("Олеся_нет", "Олеся_нет", "Ж", None),
])
def test_canonical_fio(s, n, sex, exp):
    got = canonical_fio(s, n, NAMES, sex)
    if exp:
        assert got[:2] == exp


def test_one_word_in_both_fields():
    s, n, notes = canonical_fio("Мария", "Мария", NAMES, "Ж")
    assert "одно слово в обоих полях" in notes


def _card(i, s, n, bd, nl=1, nr=0, sex="М", contacts=()):
    return {"id": i, "s0": s, "n0": n, "s": s, "n": n, "bd": bd, "sex": sex, "nl": nl, "nr": nr,
            "contacts": set(contacts), "notes": [], "one_word": False}


def test_groups_and_resolution():
    cards = {c["id"]: c for c in [
        _card(1, "Буров", "Андрей", "1990-01-01", nl=5, nr=2, contacts=["b@y.ru"]),
        _card(2, "Андрей", "Буров", "1990-01-01"),                              # D2
        _card(3, "Буров", "Андрей", "1900-01-01", contacts=["b@y.ru"]),         # заглушка + общий контакт
        _card(4, "Буров", "Андрей", "1990-01-11"),                              # одна цифра
        _card(5, "Иванов", "Мария", "1985-05-05", sex="Ж", nl=3),
        _card(6, "Иванова", "Мария", "1985-05-05", sex="Ж"),                    # -а
        _card(7, "Буров", "Андрей", "1960-03-03", contacts=["x@y.ru"]),         # другой человек, контакта нет
    ]}
    groups = {frozenset(ids): cats for ids, cats in find_groups(cards)}
    assert frozenset({1, 2, 3, 4}) in groups
    assert frozenset({5, 6}) in groups
    assert not any(7 in g for g in groups)
    assert resolve_group({1, 2, 3, 4}, cards, NAMES) == (1, "Буров", "Андрей", "1990-01-01")
    assert resolve_group({5, 6}, cards, NAMES)[1] == "Иванова"


def test_foreign_multiword_kept():
    assert canonical_fio("Fakhry", "Sherif Ashraf", NAMES, "М")[:2] == ("Fakhry", "Sherif Ashraf")


def test_make_decisions(tmp_path):
    import openpyxl
    from openpyxl.styles import PatternFill
    from scripts.clean_clients import make_decisions
    gh = ["Группа", "Что совпало", "id", "Выживает", "Было фамилия", "Было имя", "Было ДР", "Станет фамилия",
          "Станет имя", "Станет ДР", "Пол", "Заявок", "Результатов"]
    fh = ["id", "Было фамилия", "Было имя", "Станет фамилия", "Станет имя", "ДР", "Пол", "Заявок", "Результатов", "Что сделано"]

    def book(path, groups, latin, red=None):
        wb = openpyxl.Workbook()
        for title_, head, rows in (("Латиница", fh, latin), ("ФИО", fh, []), ("Склейка", gh, groups),
                                   ("Склейка ДР ±15 лет", gh, []), ("Не решено", fh, [[9, "Павел", "Павел"]])):
            ws = wb.create_sheet(title_)
            ws.append(head)
            for r in rows:
                ws.append(r)
                if red and r[0] == red:
                    ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="FFFF0000")
        wb.save(path)

    orig = [[1, "", 10, "да", "", "", "1984-01-01", "Иванов", "Иван", "2009-05-09"],
            [1, "", 11, None, "", "", "2009-05-09", "Иванов", "Иван", "2009-05-09"],
            [1, "", 12, None, "", "", "2009-05-09", "Иванов", "Иван", "2009-05-09"],
            [2, "", 20, "да", "", "", "1990-01-01", "Салников", "Родион", "1990-01-01"],
            [2, "", 21, None, "", "", "1990-01-01", "Салников", "Родион", "1990-01-01"]]
    latin = [[21, "Salnikov", "Rodion", "Салников", "Родион"], [30, "Weber", "Aleksandr", "Вебер", "Александр"],
             [31, "Test", "Runner", "Test", "Runner"]]
    user = [list(r) for r in orig]
    user[0][9] = "1984-01-01"                  # 10 — отдельный человек (родитель)
    user[1][3] = "да"                          # 11 — ребёнок, к нему 12
    user_latin = [list(r) for r in latin]
    user_latin[0][3] = "Сальников"
    book(tmp_path / "o.xlsx", orig, latin)
    book(tmp_path / "u.xlsx", user, user_latin, red=31)
    d = make_decisions(tmp_path / "u.xlsx", tmp_path / "o.xlsx")
    assert not d["errors"]
    g = {x["survivor"]: x for x in d["groups"]}
    assert g[10]["members"] == [10] and g[10]["birthday"] == "1984-01-01"
    assert sorted(g[11]["members"]) == [11, 12]
    assert g[20]["surname"] == "Сальников"     # правка на листе «Латиница» важнее группы
    assert d["renames"] == {30: ("Вебер", "Александр")}
    assert d["exclude"] == [9, 31]
