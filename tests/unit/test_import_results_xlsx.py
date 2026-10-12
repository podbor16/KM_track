import pytest

import openpyxl

from scripts.import_results_xlsx import finish_columns, parse, rank, same_card_collisions, set_bibs, sportmaniacs_rows


def test_three_finish_columns_clean_official_pace():
    h = ["Bib", "Finish", "Finish", "Finish"]                     # чистое / официальное / темп (Ночной 2025)
    rows = [(1, "00:16:01", "00:16:02", "3'12\"/km"), (2, "00:20:00", "00:20:05", "4'00\"/km")]
    assert finish_columns(h, rows) == (1, 2)


def test_named_clean_column_and_handicap():
    h = ["Bib", "Start", "Finish", "чистое "]                      # Снежная: официальное — от первого выстрела
    rows = [(1, "00:09:46", "00:38:50", "00:29:03")]
    assert finish_columns([x.strip() for x in h], rows) == (3, 2)


def test_single_finish_column_is_official_only():
    assert finish_columns(["Bib", "Результат"], [(1, "00:15:11")]) == (None, 1)


def test_clean_greater_than_official_is_error():
    h = ["Bib", "Finish \"чистое\"", "Finish"]
    rows = [(1, "00:30:00", "00:29:00")]
    with pytest.raises(ValueError):
        finish_columns(h, rows)



def test_group_headers_by_age(tmp_path):
    # Весна 2024: категория — строка-заголовок группы над участниками, пол — из неё же
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "Номер", "Фамилия", "Имя", "Дата рождения", "Статус", "Результат"])
    ws.append(["Мужчины до 49 лет (1975 г.р. и младше)"])
    ws.append([1, 1003, "Тарасов", "Павел", "05/08/2001", "Финишировал", "00:15:11"])
    ws.append([])
    ws.append(["Женщины 60 лет и старше (1964 г.р. и старше)"])
    ws.append([1, 55, "Иванова", "Анна", "01/02/1960", "Финишировал", "00:30:00"])
    path = tmp_path / "p.xlsx"
    wb.save(path)
    rows = parse(path, 2024)
    assert [(r["category"], r["sex"]) for r in rows] == [("М49", "Мужчина"), ("Ж60+", "Женщина")]
    # одна колонка времени — только официальное, чистого нет
    assert [(r["gun"], r["clean"]) for r in rows] == [(911, None), (1800, None)]
    rank(rows)
    assert (rows[0]["rank_absolute"], rows[0].get("rank_absolute_clean")) == (1, None)



def test_pacer_group_has_no_category(tmp_path):
    # Жара 2024: группа «Пейсмейкеры» в конце — не наследует предыдущую группу
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "Номер", "Фамилия", "Имя", "Год рождения", "Статус", "Результат"])
    ws.append(["Женщины 1949 г. р. и старше"])
    ws.append([1, 55, "Антипина", "Лидия", 1949, "Финишировал", "02:30:00"])
    ws.append(["Пейсмейкеры"])
    ws.append([1, 900, "Глазунов", "Владимир", None, "Финишировал", "01:45:00"])
    path = tmp_path / "p.xlsx"
    wb.save(path)
    rows = parse(path, 2024)
    assert [(r["category"], r["sex"], r["pacer_group"]) for r in rows] == [("Ж75+", "Женщина", False), ("", "", True)]



def test_set_bib():
    rows = [{"surname": "Давыденко", "name": "Петр", "bib": "515", "start_number": 515},
            {"surname": "Трегубов", "name": "Сергей", "bib": "515", "start_number": 515}]
    set_bibs(rows, ["Давыденко Пётр=593"])
    assert [r["start_number"] for r in rows] == [593, 515] and rows[0]["bib"] == "593"
    with pytest.raises(ValueError):
        set_bibs(rows, ["Иванов Иван=1"])



def test_russian_sex_and_category_headers(tmp_path):
    # Ночной 2023: «Пол» (Male/Female); Снежная 2023: «Категория» и единственное «Чистое время» — официальное
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["#", "Номер", "Имя", "Фамилия", "Статус", "Пол", "Категория", "Чистое время", "Дата рождения"])
    ws.append([1, 11, "Валентин", "Тяпкин", "Финишировал", "Male", "Мужчины", "00:31:28", "30/03/1957"])
    path = tmp_path / "p.xlsx"
    wb.save(path)
    r = parse(path, 2023)[0]
    assert (r["sex"], r["category"], r["gun"], r["clean"]) == ("Мужчина", "Мужчины", 1888, None)



def test_full_name_column_without_status(tmp_path):
    # Жара 2023: «Имя» = «Имя Фамилия», статуса нет (только финишировавшие), год рождения в «Date of Birth»
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ИТОГОВЫЙ ПРОТОКОЛ РЕЗУЛЬТАТОВ"])
    ws.append(["#", "Номер", "Имя", "Результат", "Клуб", "Город", "Date of Birth"])
    ws.append(["женщины 2010−2011 г. р."])
    ws.append(["1", "1179", "Кира Огер", "00:48:05", "Сибиряк", "Красноярск", 2011])
    ws.append(["Вне зачета"])
    ws.append(["1", "777", "Александр Ермолов", "02:19:01", None, "Красноярск", "-"])
    ws.append([])
    ws.append(["Главный судья", "Кондоба А.С."])
    path = tmp_path / "p.xlsx"
    wb.save(path)
    r = parse(path, 2023)[0]
    assert (r["surname"], r["name"], r["race_status"], r["gun"], r["birth_year"], r["birthday"], r["sex"], r["start_number"]) == \
        ("Огер", "Кира", "Finished", 2885, 2011, "1900-01-01", "Женщина", 1179)
    assert r["category"] == "Ж12-13"
    rows = parse(path, 2023)
    assert len(rows) == 2 and (rows[1]["surname"], rows[1]["category"], rows[1]["pacer_group"]) == ("Ермолов", "", True)



def test_same_card_collisions():
    # Жара 2023: две Анисимовы Елены 1985 с результатом, ДР — заглушка → одна карточка
    row = lambda bib, status, bd="1985-01-01": {"surname": "Анисимова", "name": "Елена", "birthday": bd,
                                                "start_number": bib, "race_status": status}
    rows = [row(867, "Finished"), row(945, "Finished"), row(1, "Not started"), row(2, "Finished", "1985-12-06")]
    assert [[r["start_number"] for r in g] for g in same_card_collisions(rows)] == [[867, 945]]



def _xlsx(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    path = tmp_path / "p.xlsx"
    wb.save(path)
    return path


def test_kids_year_groups_above_header(tmp_path):
    # Детский 2022: «2016 ГОД» над заголовками, дальше группы — просто год
    path = _xlsx(tmp_path, [["1 КМ"], ["2016 ГОД"], ["#", "Номер", "Имя", "Фамилия", "Пол", "Результат", "Город"],
                            ["1", "601", "Тимофей", "Мохов", "Male", "00:04:24", "Канск"],
                            [2014], ["1", "401", "Анна", "Иванова", "Female", "00:04:50", "Канск"],
                            [], ["Главный судья ", "Ситников В.Н."]])
    rows = parse(path, 2022)
    assert [(r["category"], r["birth_year"]) for r in rows] == [("Мальчики 2016 г.р.", 2016), ("Девочки 2014 г.р.", 2014)]


def test_dns_in_result_column_and_retired(tmp_path):
    # Х Трейл 2022: статуса нет, «DNS» вместо времени; Весна 2022: «Retired»; Ночной 2022: пол «Мужчины»
    path = _xlsx(tmp_path, [["#", "Номер", "Фамилия", "Имя", "пол", "Результат"],
                            ["1", "518", "Матвеенко", "Егор", "Мужчины", "00:38:25"],
                            ["2", "7", "Иванов", "Иван", "Male", "DNS"]])
    rows = parse(path, 2022)
    assert [(r["race_status"], r["sex"]) for r in rows] == [("Finished", "Мужчина"), ("Not started", "Мужчина")]
    path = _xlsx(tmp_path, [["#", "Номер", "Имя", "Фамилия", "Статаус", "Пол", "Результат"],
                            ["1", "13", "Никита", "Калинин", "Finished", "Male", "00:16:32"],
                            ["2", "14", "Иван", "Петров", "Retired", "Male", None]])
    assert [r["race_status"] for r in parse(path, 2022)] == ["Finished", "Withdrawn"]



def test_kids_two_year_groups_and_not_started(tmp_path):
    # Детский 2021: «девочки 2011-2010», «Not started» вместо времени
    path = _xlsx(tmp_path, [["Место", "Номер", "Фамилия", "Имя", "Результат"], ["девочки 2011-2010"],
                            ["1", "180", "Шмыгина", "Анна", "00:03:44"], ["", "503", "Воронина", "Нелли", "Not started"]])
    rows = parse(path, 2021)
    assert [(r["category"], r["race_status"]) for r in rows] == [("Девочки 2010-2011 г.р.", "Finished"),
                                                                  ("Девочки 2010-2011 г.р.", "Not started")]



def test_sportmaniacs_rows():
    data = {"Rankings": [
        {"dorsal": "86", "name": "ВАЛЕНТИН ТЯПКИН", "gender": "gender_0", "category": "Мужчины 12 лет и старше",
         "officialTime": "00:38:03", "realTime": "00:28:12", "finishedRace": True},
        {"dorsal": "401", "name": "Елена Иванова ", "gender": "gender_1", "category": "All Ages",
         "officialTime": "00:00:00", "realTime": "00:00:00", "finishedRace": False},
    ]}
    a, b = sportmaniacs_rows(data, 2020)
    assert (a["surname"], a["name"], a["sex"], a["gun"], a["clean"], a["race_status"], a["start_number"]) == \
        ("Тяпкин", "Валентин", "Мужчина", 2283, 1692, "Finished", 86)
    assert (b["surname"], b["race_status"], b["gun"], b["category"]) == ("Иванова", "Not started", None, "")
