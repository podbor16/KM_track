"""src/analytics/dash_export.py — разбор ответов DataLens и сборка Excel без сети."""
import io
from datetime import datetime
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from src.analytics import dash_export as de

ENTRY = {"entryId": "dash1", "data": {"tabs": [{
    "id": "XbA", "title": "Динамика регистраций",
    "layout": [{"i": "w2", "x": 0, "y": 5}, {"i": "w1", "x": 0, "y": 2}, {"i": "c1", "x": 0, "y": 0}],
    "items": [
        {"id": "w2", "type": "widget", "data": {"tabs": [{"chartId": "g1", "title": "Д1 - Итог"}]}},
        {"id": "w1", "type": "widget", "data": {"tabs": [{"chartId": "m1", "title": "Д4 - Остаток"},
                                                          {"chartId": "t1", "title": "Д3 - Продажи"}]}},
        {"id": "c1", "type": "group_control", "data": {"group": [
            {"title": "Год", "source": {"datasetFieldId": "event_year_jr76"}},
            {"title": "Мероприятие", "source": {"datasetFieldId": "event_name_2iel", "multiselectable": True}}]}},
    ]}]}}

METRIC = {"data": [{"title": "Остаток", "content": {"current": {"value": 144}}}]}
TABLE = {"data": {"head": [{"name": "event_name"}, {"id": "id"}],
                  "rows": [{"cells": [{"value": "Жара"}, {"value": 997}]}]}}
GRAPH = {"data": {"xAxis": {"categories": ["01. 40+", "02. 24-39"]}, "series": {"data": [
    {"name": 2024, "type": "line", "data": [{"x": 0, "y": 48}, {"x": 1, "y": 71}]},
    {"name": "2025", "type": "line", "data": [{"x": 1, "y": 157}, {"x": 9, "y": 1}]}]}}}


@pytest.fixture(autouse=True)
def entry():
    with patch.dict(de._entry_cache, {"ts": 1e18, "entry": ENTRY}):
        yield


def test_normalize_kinds():
    assert de.normalize(METRIC) == {"kind": "metric", "label": "Остаток", "value": 144}
    assert de.normalize(TABLE) == {"kind": "table", "columns": ["event_name", "id"], "rows": [["Жара", 997]]}
    g = de.normalize(GRAPH)
    assert g["chart"] == "line" and g["categories"] == ["01. 40+", "02. 24-39"]
    assert g["series"] == [{"name": "2024", "type": "line", "values": [48, 71]},
                           {"name": "2025", "type": "line", "values": [None, 157]}]
    bar = {"data": {"xAxis": {"categories": ["a"]}, "series": {"data": [{"name": "x", "type": "bar-x", "data": []}]}}}
    assert de.normalize(bar)["chart"] == "bar"
    assert de.normalize({"data": "???"}) is None


def test_charts_in_layout_order_and_filter_titles():
    assert de.charts("XbA") == [("m1", "Д4 - Остаток"), ("t1", "Д3 - Продажи"), ("g1", "Д1 - Итог")]
    assert de._filter_titles("XbA") == {"event_year_jr76": "Год", "event_name_2iel": "Мероприятие"}
    assert de.tabs() == [{"id": "XbA", "title": "Динамика регистраций", "export": True}]
    with pytest.raises(ValueError):
        de.charts("nope")


def test_lead_field_and_sheet_title():
    assert de._lead_field("event_name_2iel") == "event_name"
    assert de._lead_field("d5ceaa10-4923") is None and de._lead_field(None) is None
    used = set()
    assert de._sheet_title("Итог: [а/б]", used) == "Итог   а б"
    long = "x" * 40
    assert de._sheet_title(long, used) == "x" * 31 and de._sheet_title(long, used) == "x" * 28 + " 2"


def test_build_xlsx_sheets_charts_and_params():
    sent = []

    def run(chart_id, params, tab_id):
        sent.append((chart_id, params, tab_id))
        return {"m1": METRIC, "t1": TABLE, "g1": GRAPH}[chart_id]

    with patch.object(de, "run_chart", side_effect=run):
        data, name = de.build_xlsx("XbA", {"event_name_2iel": ["Жара", "Весна"]}, now=datetime(2026, 10, 26, 12, 5))
    assert name == "analytics_Динамика_регистраций_20261026-1205.xlsx"
    # селектор без значения уходит в DataLens явно пустым («все»)
    assert all(p == {"event_year_jr76": "", "event_name_2iel": ["Жара", "Весна"]} and t == "XbA" for _, p, t in sent)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Сводка", "Д3 - Продажи", "Д1 - Итог"]
    rows = list(wb["Сводка"].iter_rows(values_only=True))
    assert ("Год", "все", None) in rows and ("Мероприятие", "Жара, Весна", None) in rows
    assert ("Д4 - Остаток", "Остаток", 144) in rows
    assert list(wb["Д3 - Продажи"].iter_rows(min_row=2, values_only=True)) == [("event_name", "id"), ("Жара", 997)]
    g = wb["Д1 - Итог"]
    assert list(g.iter_rows(min_row=2, max_col=3, values_only=True)) == [
        (None, "2024", "2025"), ("01. 40+", 48, None), ("02. 24-39", 71, 157)]
    assert len(g._charts) == 1


def test_build_xlsx_rejects_tab_not_enabled():
    with pytest.raises(ValueError):
        de.build_xlsx("nA", {})


def test_call_sends_embed_headers():
    captured = {}

    class Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            pass

    def urlopen(req, timeout):
        captured.update(url=req.full_url, method=req.get_method(), headers=dict(req.header_items()))
        return Resp(b'{"ok": 1}')

    with patch.object(de, "_embed", return_value={"id": "e1"}), patch.object(de, "make_embed_token", return_value="TOK"), \
            patch.object(de.urllib.request, "urlopen", side_effect=urlopen):
        assert de.run_chart("g1", {"a": ""}, "XbA") == {"ok": 1}
    h = {k.lower(): v for k, v in captured["headers"].items()}
    assert captured["method"] == "POST" and captured["url"].endswith("/charts/embeds/api/run")
    assert h["x-dl-embed-token"] == "TOK" and "Mozilla" in h["user-agent"]
    assert h["x-dash-info"] == "dashId=dash1&dashTabId=XbA"


def test_call_retries_once_on_transient_error():
    import urllib.error

    class Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            pass

    def err(code):
        return urllib.error.HTTPError("u", code, "x", {}, io.BytesIO())

    with patch.object(de, "_embed", return_value={"id": "e1"}), patch.object(de, "make_embed_token", return_value="T"), \
            patch.object(de, "RETRY_DELAY", 0):
        with patch.object(de.urllib.request, "urlopen", side_effect=[err(424), Resp(b"{}")]) as u:
            assert de.run_chart("g1", {}, "XbA") == {} and u.call_count == 2
        with patch.object(de.urllib.request, "urlopen", side_effect=[err(424), err(424)]), pytest.raises(urllib.error.HTTPError):
            de.run_chart("g1", {}, "XbA")
        with patch.object(de.urllib.request, "urlopen", side_effect=[err(400)]) as u, pytest.raises(urllib.error.HTTPError):
            de.run_chart("g1", {}, "XbA")
        assert u.call_count == 1
