#!/usr/bin/env python3
"""
Чистка карточек клиентов (2026-09-26, решения пользователя по аудиту):

1. ФИО в полях: невидимые символы, лишние пробелы, латиница-двойники и
   посторонние символы; регистр «Иванов»; отчество и инициалы убираются;
   ФИО целиком в одном/обоих полях и имя/фамилия местами — раскладываются по
   полям; латиница -> кириллица, если это русское ФИ (имя узнаётся по
   частотным именам с учётом пола), иностранные ФИ остаются латиницей.
2. Склейка дублей (все категории аудита D1–D8): те же ФИ+ДР, ФИ переставлены,
   ДР-заглушка, ДР отличается на цифру/день↔месяц, разные ДР при общем
   email/телефоне, опечатка в букве фамилии/имени, -а в фамилии; одно слово
   вместо ФИ — карточка с тем же словом, той же ДР и общим контактом.
   Выживает карточка с наибольшим числом результатов, затем заявок; ФИО и ДР
   группы — взвешенное большинство (вес — результаты и заявки), -а в фамилии —
   по полу. Заявки и результаты остальных карточек переносятся к ней.
3. ФИО и ДР в leads/results затронутых карточек приводятся к карточке.
   Флаги заявок (is_duplicate, is_new) не пересчитываются: is_duplicate=1
   убрал бы обе заявки из выгрузки стартового списка.

  python scripts/clean_clients.py --report /tmp/review.xlsx                              # dry-run
  python scripts/clean_clients.py --review /tmp/review.xlsx --report /tmp/r2.xlsx --apply --backup /root/backups/x.json

--review — проверенный пользователем файл: колонки «Станет фамилия/имя»
листов «Латиница» и «ФИО» заменяют вычисленные.
"""

import argparse
import collections
import itertools
import json
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).parent.parent))

SENTINELS = {"1900-01-01", "1905-01-01"}

# ---------------------------------------------------------------- текст ФИО

_INVISIBLE = re.compile("[​-‏⁠﻿­̀́]")
_HOMOGLYPH = str.maketrans("aAeEoOpPcCxXyYkKmMtTHBh", "аАеЕоОрРсСхХуУкКмМтТНВһ")
_TO_LATIN = str.maketrans("аАеЕоОрРсСхХуУкКмМтТНВ", "aAeEoOpPcCxXyYkKmMtTHB")
_CYR = re.compile(r"[а-яё]", re.I)
_LAT = re.compile(r"[a-z]", re.I)
_PATR_FEMALE = re.compile(r"(вна|ична)$", re.I)     # фамилий с такими окончаниями почти нет
_PATR_MALE = re.compile(r"вич$", re.I)               # «Валисевич», «Богданкевич» — фамилии

_TRANSLIT = [
    ("shch", "щ"), ("sch", "щ"), ("iia", "ия"), ("zh", "ж"), ("kh", "х"), ("ts", "ц"), ("tz", "ц"),
    ("ch", "ч"), ("sh", "ш"), ("yu", "ю"), ("iu", "ю"), ("ju", "ю"), ("ya", "я"), ("ja", "я"),
    ("ia", "ия"), ("yo", "ё"), ("jo", "ё"), ("ye", "е"), ("je", "е"), ("ii", "ий"), ("iy", "ий"),
    ("yi", "ый"), ("ey", "ей"), ("ay", "ай"), ("oy", "ой"), ("uy", "уй"), ("ph", "ф"), ("x", "кс"),
    ("a", "а"), ("b", "б"), ("v", "в"), ("w", "в"), ("g", "г"), ("d", "д"), ("e", "е"), ("z", "з"),
    ("i", "и"), ("j", "й"), ("k", "к"), ("l", "л"), ("m", "м"), ("n", "н"), ("o", "о"), ("p", "п"),
    ("r", "р"), ("s", "с"), ("t", "т"), ("u", "у"), ("f", "ф"), ("h", "х"), ("c", "к"), ("q", "к"),
]


def norm(s):
    """Ключ сравнения: без регистра, ё=е, без лишних пробелов (как collation БД)."""
    return re.sub(r"\s+", " ", (s or "").strip().lower().replace("ё", "е"))


def translit(word):
    w, out, i = word.lower(), [], 0
    while i < len(w):
        if w[i] in "aeou" and w[i + 1:i + 2] == "y" and w[i + 2:i + 3] in tuple("aeiou"):
            out.append(dict(zip("aeou", "аеоу"))[w[i]])   # «Troyakova» -> «Троякова»
            i += 1
            continue
        if w[i] == "y" and not w.startswith(("ya", "yu", "yo", "ye", "yi"), i):
            # конечная y после согласной — «ий» (Dmitry), иначе «ы»
            out.append("ий" if i == len(w) - 1 and i and w[i - 1] not in "aeiou" else "ы")
            i += 1
            continue
        for lat, cyr in _TRANSLIT:
            if w.startswith(lat, i):
                out.append(cyr)
                i += len(lat)
                break
        else:
            out.append(w[i])
            i += 1
    return "".join(out)


def title(word):
    """«ИВАНОВ»/«иванов» -> «Иванов» (по частям через дефис); «МакКой» не трогаем."""
    if not (word.isupper() or word.islower() or word[:1].islower()):
        return word
    return "-".join(p[:1].upper() + p[1:].lower() for p in word.split("-"))


def lev(a, b, limit=2):
    """Расстояние Дамерау (перестановка соседних букв — одна правка: «Adnrei»)."""
    if abs(len(a) - len(b)) > limit:
        return limit + 1
    d = [[i + j if not i * j else 0 for j in range(len(b) + 1)] for i in range(len(a) + 1)]
    for i in range(1, len(a) + 1):
        for j in range(1, len(b) + 1):
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + (a[i - 1] != b[j - 1]))
            if i > 1 and j > 1 and a[i - 1] == b[j - 2] and a[i - 2] == b[j - 1]:
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 1)
    return d[-1][-1]


def tokens(field):
    s = _INVISIBLE.sub("", field or "").replace(" ", " ")
    s = re.sub(r"\s*-\s*", "-", s)                     # «Петров- Дельверс»
    s = re.sub(r"[^\w\s\-'’]|[\d_]", " ", s)          # «Салимжанов.», «✅», «Ксения69_»
    out = []
    for t in s.split():
        t = t.strip("-'’")
        if _CYR.search(t) and _LAT.search(t):
            if len(_LAT.findall(t)) > len(_CYR.findall(t)):  # «Моrotskiy» — латиница с кириллическими М, о
                t = translit(t.translate(_TO_LATIN))
            else:                                            # «Нинa», «ЛаZученко», «Светланd»
                t = "".join(ch if not _LAT.match(ch) else
                            ch.translate(_HOMOGLYPH) if ch.translate(_HOMOGLYPH) != ch else translit(ch)
                            for ch in t)
        if t:
            out.append(t)
    return out


class Names:
    """Частотные имена (с полом) и фамилии — по заявкам и результатам."""

    def __init__(self, rows):
        name_f, surn_f = collections.Counter(), collections.Counter()
        sex = collections.defaultdict(collections.Counter)
        for s, n, x in rows:
            for t in tokens(n)[:1]:
                name_f[norm(t)] += 1
                sex[norm(t)][(x or "")[:1].upper()] += 1
            for t in tokens(s)[:1]:
                surn_f[norm(t)] += 1
        self.first = {n: k for n, k in name_f.items() if k >= 5 and surn_f[n] < k / 3 and _CYR.search(n)}
        self.sex = {n: sex[n].most_common(1)[0][0] for n in self.first}

    def is_first(self, t):
        return norm(t) in self.first

    def match_first(self, lat, sex=""):
        """Латинское имя -> частотное русское имя (или None). Точное совпадение
        транслита (в т.ч. с «ь»: Igor -> Игорь) — без учёта пола; близкое —
        только того же пола («Petra» не «Петр»)."""
        c = translit(lat)
        for exact in (c, c + "ь"):
            if exact in self.first:
                return exact
        limit = 0 if len(c) <= 4 else 1 if len(c) == 5 else 2
        best = None
        for n, k in self.first.items():
            d = lev(n, c, limit) if limit else 1
            if d > limit or (sex and self.sex[n] in "МЖ" and self.sex[n] != sex[:1].upper()):
                continue
            if best is None or (d, -k) < best[0]:
                best = ((d, -k), n)
        return best[1] if best else None


def canonical_fio(surname, name, names, sex=""):
    """-> (фамилия, имя, заметки). Пустая строка — поле определить не удалось."""
    ts, tn = tokens(surname), tokens(name)
    if norm(surname) == norm(name):
        if len(ts) < 2:                                 # «Олеся» «Олеся» — второго слова нет
            w = title(ts[0]) if ts else ""
            return w, w, ["одно слово в обоих полях"]
        tn = []                                         # ФИО целиком в обоих полях
    notes = []
    # инициалы: «С.», «К.», «ЕП»
    initial = lambda t: len(t) == 1 or (len(t) == 2 and t.isupper())
    if any(initial(t) for t in ts + tn):
        notes.append("инициал убран")
    ts, tn = [t for t in ts if not initial(t)], [t for t in tn if not initial(t)]
    # отчество: -вна/-ична — всегда; -вич — сразу после имени и если есть другая фамилия
    # («Наталия Валисевич», «Богданкевич Ратибор» — фамилии)
    other_surname = lambda x: any(u != x and not names.is_first(u) and not _PATR_MALE.search(u)
                                  and not _PATR_FEMALE.search(u) for u in ts + tn)
    kept = []
    for field in (ts, tn):
        out = []
        for k, t in enumerate(field):
            male = _PATR_MALE.search(t) and k > 0 and names.is_first(field[k - 1]) and other_surname(t)
            if len(t) > 5 and not names.is_first(t) and (_PATR_FEMALE.search(t) or male):
                notes.append("отчество убрано")
            else:
                out.append(t)
        kept.append(out)
    ts, tn = kept
    words = [("s", t) for t in ts] + [("n", t) for t in tn]

    # латиница: русское ФИ -> кириллица (транслит), иностранное остаётся
    if any(not _CYR.search(t) for _, t in words):
        cyr = [t for _, t in words if _CYR.search(t)]
        if any(names.is_first(t) for t in cyr) and any(not names.is_first(t) for t in cyr):
            words = [x for x in words if _CYR.search(x[1])]          # «Прусаков … Prusakov Vladimir»
            notes.append("латинский дубль ФИО убран")
        else:
            # имя ищем сначала в поле имени; в поле фамилии — только если в поле имени его нет
            fm = {}
            for field in ("n", "s"):
                if field == "s" and (any(names.is_first(t) for f, t in words if f == "n" and _CYR.search(t))
                                     or any(fm.values())):
                    break
                for f, t in words:
                    if f == field and not _CYR.search(t):
                        fm[t] = names.match_first(t, sex)
            if cyr or any(fm.values()):
                words = [(f, t if _CYR.search(t) else fm.get(t) or title(translit(t))) for f, t in words]
                notes.append("латиница -> кириллица")
            else:
                notes.append("иностранное ФИ")

    # имя: частотное имя из поля имени, иначе из поля фамилии, иначе поле имени
    firsts = [x for x in words if names.is_first(x[1])]
    pick = (next((x for x in firsts if x[0] == "n"), None) or (firsts[0] if firsts else None)
            or next((x for x in words if x[0] == "n"), None))
    if pick and not names.is_first(pick[1]) and len([f for f, _ in words if f == "n"]) > 1:
        return "", "", notes + ["не удалось разобрать имя"]     # «Алмазная Крошка»
    new_name = title(pick[1]) if pick else ""
    if "-" in new_name:                                          # «Есения-Принцесса» -> «Есения»
        parts = new_name.split("-")
        if any(names.is_first(p) for p in parts) and not all(names.is_first(p) for p in parts):
            new_name = "-".join(p for p in parts if names.is_first(p))
    left = [x for x in words if x is not pick]
    # фамилия: не-имена, сначала из поля фамилии; лишние имена (родителя) отбрасываются
    cand = [x for x in left if not names.is_first(x[1])] or left
    parts = [t for f, t in cand if f == "s"] or [t for f, t in cand]
    if pick and pick[0] == "s" and cand and all(f == "n" for f, _ in cand):
        notes.append("имя и фамилия переставлены")
    uniq = []
    for t in parts:
        if norm(t) not in {norm(u) for u in uniq} and (not pick or norm(t) != norm(pick[1])):
            uniq.append(t)
    if len(words) > 2 or len(uniq) + (1 if pick else 0) < len(words):
        notes.append("лишние слова убраны")
    if not uniq:
        notes.append("нет фамилии")
    return " ".join(title(t) for t in uniq), new_name, notes


# ---------------------------------------------------------------- дубли

def _phone(v):
    d = re.sub(r"\D", "", v or "")
    return d[-10:] if len(d) >= 10 else ""


def _email(v):
    v = (v or "").strip().lower()
    return "" if v in ("", "example@mail.ru") else v


def _lev1(a, b):
    return a != b and lev(a, b, 1) == 1


def _near_bd(a, b):
    if a in SENTINELS or b in SENTINELS or a == b:
        return False
    (ya, ma, da), (yb, mb, db) = a.split("-"), b.split("-")
    return (ya, ma, da) == (yb, db, mb) or sum(x != y for x, y in zip(a, b)) == 1


def find_groups(cards):
    """cards: {id: {s, n, bd, contacts, ...}} с каноническими s/n.
    -> [(set(id), {категории})] — группы из ≥2 карточек."""
    parent = {i: i for i in cards}
    cats = collections.defaultdict(set)

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b, cat):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra
        cats[(min(a, b), max(a, b))].add(cat)

    by_fi, by_nb, by_sb, by_word = (collections.defaultdict(list) for _ in range(4))
    for i, c in cards.items():
        fi = (norm(c["s"]), norm(c["n"]))
        by_fi[fi].append(i)
        if c["bd"] not in SENTINELS:
            by_nb[(fi[1], c["bd"])].append(i)
            by_sb[(fi[0], c["bd"])].append(i)
    for fi, ids in by_fi.items():
        for a, b in itertools.combinations(ids, 2):
            ca, cb = cards[a], cards[b]
            shared = ca["contacts"] & cb["contacts"]
            if ca["bd"] == cb["bd"]:
                union(a, b, "D1 те же ФИ и ДР")
            elif ca["bd"] in SENTINELS or cb["bd"] in SENTINELS:
                real = {cards[x]["bd"] for x in ids if cards[x]["bd"] not in SENTINELS}
                if shared or len(real) == 1:
                    union(a, b, "D3a/D4 ДР-заглушка")
            elif shared:
                union(a, b, "D3b разные ДР, общий контакт")
            elif _near_bd(ca["bd"], cb["bd"]):
                union(a, b, "D5 ДР отличается на цифру")
        swapped = (fi[1], fi[0])
        if swapped != fi:
            for a in ids:
                for b in by_fi.get(swapped, []):
                    if cards[a]["bd"] == cards[b]["bd"]:
                        union(a, b, "D2 ФИ переставлены")
    for group_by, cat_a, cat_b in ((by_nb, "D6 опечатка в фамилии", "D6b -а в фамилии"), (by_sb, "D7 опечатка в имени", None)):
        for (_, _), ids in group_by.items():
            for a, b in itertools.combinations(ids, 2):
                x, y = ((norm(cards[a]["s"]), norm(cards[b]["s"])) if cat_b else (norm(cards[a]["n"]), norm(cards[b]["n"])))
                if _lev1(x, y):
                    union(a, b, cat_b if cat_b and (x + "а" == y or y + "а" == x) else cat_a)
    # одно слово вместо ФИ: карточка с этим словом в фамилии или имени, той же ДР и общим контактом
    for i, c in cards.items():
        if c.get("one_word"):
            w = norm(c["s"] or c["n"])
            for j, d in cards.items():
                if j != i and d["bd"] == c["bd"] and c["contacts"] & d["contacts"] and \
                        w in (norm(d["s"]), norm(d["n"])) and not d.get("one_word"):
                    union(i, j, "L одно слово вместо ФИ")
    groups = collections.defaultdict(set)
    for i in cards:
        groups[find(i)].add(i)
    out = []
    for ids in groups.values():
        if len(ids) > 1:
            gc = set()
            for a, b in itertools.combinations(sorted(ids), 2):
                gc |= cats.get((a, b), set())
            out.append((ids, gc))
    return out


def resolve_group(ids, cards, names):
    """-> (выжившая id, фамилия, имя, ДР) группы."""
    w = {i: 10 * cards[i]["nr"] + cards[i]["nl"] + 1 for i in ids}
    survivor = max(ids, key=lambda i: (cards[i]["nr"], cards[i]["nl"], cards[i]["bd"] not in SENTINELS, -i))

    def vote(field, ok=lambda v: True):
        tally = collections.Counter()
        for i in ids:
            v = cards[i][field]
            if v and ok(v):
                tally[v] += w[i]
        return tally.most_common(1)[0][0] if tally else cards[survivor][field]

    cyr = lambda v: bool(_CYR.search(v))
    has_cyr = any(cyr(cards[i]["n"]) for i in ids)
    name = vote("n", lambda v: (cyr(v) or not has_cyr) and (names.is_first(v) or not any(names.is_first(cards[i]["n"]) for i in ids)))
    surname = vote("s", lambda v: (cyr(v) or not has_cyr) and len(v) > 1)
    # -а в фамилии — по полу группы
    sexes = collections.Counter(cards[i]["sex"] for i in ids if cards[i]["sex"])
    sex = sexes.most_common(1)[0][0] if sexes else ""
    variants = {cards[i]["s"] for i in ids}
    base = surname[:-1] if surname.endswith("а") and surname[:-1] in variants else surname
    if base + "а" in variants or base != surname:
        surname = base + "а" if sex == "Ж" else base
    bd = vote("bd", lambda v: v not in SENTINELS)
    return survivor, surname, name, bd


# ---------------------------------------------------------------- БД

def load(conn):
    cur = conn.cursor(dictionary=True)
    cur.execute("""
        SELECT c.id, c.surname, c.name, c.birthday, c.email, c.phone,
               (SELECT COUNT(*) FROM leads l WHERE l.client_id = c.id) nl,
               (SELECT COUNT(*) FROM results r WHERE r.client_id = c.id) nr
        FROM clients c WHERE c.id <> 0""")
    clients = cur.fetchall()
    cur.execute("SELECT client_id, surname, name, sex, email, phone FROM leads")
    leads = cur.fetchall()
    cur.execute("SELECT client_id, surname, name, sex FROM results")
    results = cur.fetchall()
    cur.close()
    return clients, leads, results


def build_cards(clients, leads, results, names, overrides):
    sex = collections.defaultdict(collections.Counter)
    contacts = collections.defaultdict(set)
    for r in leads:
        sex[r["client_id"]][(r["sex"] or "")[:1].upper()] += 1
        contacts[r["client_id"]] |= {x for x in (_email(r["email"]), _phone(r["phone"])) if x}
    for r in results:
        sex[r["client_id"]][(r["sex"] or "")[:1].upper()] += 1
    cards = {}
    for c in clients:
        x = sex[c["id"]]
        x.pop("", None)
        sx = x.most_common(1)[0][0] if x else ""
        s, n, notes = canonical_fio(c["surname"], c["name"], names, sx)
        if c["id"] in overrides:
            s, n = overrides[c["id"]]
            notes = notes + ["проверено вручную"]
        cards[c["id"]] = {
            "id": c["id"], "s0": c["surname"], "n0": c["name"], "s": s, "n": n, "notes": notes,
            "bd": str(c["birthday"]), "sex": sx, "nl": c["nl"], "nr": c["nr"],
            "contacts": contacts[c["id"]] | {x for x in (_email(c["email"]), _phone(c["phone"])) if x},
            "one_word": "одно слово в обоих полях" in notes or not s or not n,
        }
    return cards


def plan_changes(cards, names):
    groups = find_groups(cards)
    final, merged_into, group_info = {}, {}, []
    for gi, (ids, cats) in enumerate(sorted(groups, key=lambda g: min(g[0])), 1):
        survivor, s, n, bd = resolve_group(ids, cards, names)
        final[survivor] = (s, n, bd)
        for i in ids:
            if i != survivor:
                merged_into[i] = survivor
        group_info.append((gi, sorted(ids, key=lambda i: (i != survivor, i)), cats, survivor, (s, n, bd)))
    for i, c in cards.items():
        if i not in final and i not in merged_into and not c["one_word"] and (c["s"], c["n"]) != (c["s0"], c["n0"]):
            final[i] = (c["s"], c["n"], c["bd"])
    # уникальность (фамилия, имя, ДР) после изменений — как uk_client
    keys = collections.Counter()
    for i, c in cards.items():
        if i in merged_into:
            continue
        s, n, bd = final.get(i, (c["s0"], c["n0"], c["bd"]))
        keys[(norm(s), norm(n), bd)] += 1
    clash = [k for k, v in keys.items() if v > 1]
    return final, merged_into, group_info, clash


def apply(conn, cards, final, merged_into, backup):
    cur = conn.cursor(dictionary=True)
    touched = sorted(set(final) | set(merged_into))
    ph = lambda ids: ",".join(str(int(i)) for i in ids)
    snap = {}
    for chunk in (touched[i:i + 2000] for i in range(0, len(touched), 2000)):
        cur.execute(f"SELECT * FROM clients WHERE id IN ({ph(chunk)})")
        snap.setdefault("clients", []).extend(cur.fetchall())
        cur.execute(f"SELECT id, client_id, surname, name, birthday FROM leads WHERE client_id IN ({ph(chunk)})")
        snap.setdefault("leads", []).extend(cur.fetchall())
        cur.execute(f"SELECT id, client_id, surname, name, birthday FROM results WHERE client_id IN ({ph(chunk)})")
        snap.setdefault("results", []).extend(cur.fetchall())
    Path(backup).write_text(json.dumps(snap, default=str, ensure_ascii=False), encoding="utf-8")

    by_survivor = collections.defaultdict(list)
    for i, s in merged_into.items():
        by_survivor[s].append(i)
    try:
        for s, others in by_survivor.items():
            cur.execute(f"UPDATE leads SET client_id = %s WHERE client_id IN ({ph(others)})", (s,))
            cur.execute(f"UPDATE results SET client_id = %s WHERE client_id IN ({ph(others)})", (s,))
        merged = sorted(merged_into)
        for chunk in (merged[i:i + 1000] for i in range(0, len(merged), 1000)):
            cur.execute(f"DELETE FROM clients WHERE id IN ({ph(chunk)})")
        # два шага, чтобы не упереться в uk_client при взаимных переименованиях
        ids = sorted(final)
        for chunk in (ids[i:i + 1000] for i in range(0, len(ids), 1000)):
            cur.execute(f"UPDATE clients SET surname = CONCAT('~', id) WHERE id IN ({ph(chunk)})")
        for i, (s, n, bd) in final.items():
            cur.execute("UPDATE clients SET surname = %s, name = %s, birthday = %s WHERE id = %s", (s, n, bd, i))
            cur.execute("UPDATE leads SET surname = %s, name = %s, birthday = %s WHERE client_id = %s", (s, n, bd, i))
            cur.execute("UPDATE results SET surname = %s, name = %s, birthday = %s WHERE client_id = %s", (s, n, bd, i))
        for chunk in (ids[i:i + 1000] for i in range(0, len(ids), 1000)):
            recompute_aggregates(cur, chunk)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return len(snap.get("leads", [])), len(snap.get("results", []))


def recompute_aggregates(cur, ids):
    """Счётчики карточки — из фактических заявок и результатов."""
    ph = ",".join(str(int(i)) for i in ids)
    cur.execute(f"""
        UPDATE clients c
        LEFT JOIN (SELECT client_id, COUNT(*) n, COALESCE(SUM(amount), 0) s, MAX(id) mx,
                          MIN(created_at) mn, MAX(created_at) md
                   FROM leads WHERE client_id IN ({ph}) GROUP BY client_id) l ON l.client_id = c.id
        LEFT JOIN (SELECT client_id, MAX(id) r FROM results WHERE client_id IN ({ph}) GROUP BY client_id) r
               ON r.client_id = c.id
        SET c.count_leads = COALESCE(l.n, 0), c.total_amount = COALESCE(l.s, 0), c.last_lead_id = l.mx,
            c.first_lead_date = l.mn, c.last_lead_date = l.md, c.last_result_id = r.r
        WHERE c.id IN ({ph})""")
    cur.execute(f"""
        UPDATE clients c
        JOIN leads l ON l.id = (SELECT x.id FROM leads x WHERE x.client_id = c.id ORDER BY x.created_at, x.id LIMIT 1)
        SET c.first_event_id = l.event_id, c.first_event_name = l.event_name
        WHERE c.id IN ({ph})""")


# ---------------------------------------------------------------- отчёт

def read_overrides(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for sheet in ("Латиница", "ФИО"):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        h = [str(x) for x in rows[0]]
        i_id, i_s, i_n = h.index("id"), h.index("Станет фамилия"), h.index("Станет имя")
        for r in rows[1:]:
            if r[i_id]:
                out[int(r[i_id])] = (str(r[i_s] or "").strip(), str(r[i_n] or "").strip())
    return out


def write_report(path, cards, final, merged_into, group_info, clash):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Сводка"
    fill = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FB")]
    head = ["id", "Было фамилия", "Было имя", "Станет фамилия", "Станет имя", "ДР", "Пол", "Заявок", "Результатов", "Что сделано"]

    def sheet(title_, rows, widths):
        ws = wb.create_sheet(title_)
        ws.append(rows[0])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows[1:]:
            ws.append(r)
        for i, wd in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    solo = [i for i in final if i not in {g[3] for g in group_info}]
    lat = [i for i, c in cards.items() if any(n.startswith(("латиница", "иностранное")) for n in c["notes"])]
    fio = [i for i in solo if i not in lat]
    row = lambda i: [i, cards[i]["s0"], cards[i]["n0"], cards[i]["s"], cards[i]["n"], cards[i]["bd"], cards[i]["sex"],
                     cards[i]["nl"], cards[i]["nr"], ", ".join(cards[i]["notes"])]
    widths = [8, 24, 20, 24, 20, 11, 6, 8, 11, 40]
    sheet("Латиница", [head] + [row(i) for i in sorted(lat)], widths)
    sheet("ФИО", [head] + [row(i) for i in sorted(fio)], widths)
    ghead = ["Группа", "Что совпало", "id", "Выживает", "Было фамилия", "Было имя", "Было ДР", "Станет фамилия",
             "Станет имя", "Станет ДР", "Пол", "Заявок", "Результатов"]
    grows, gap = [ghead], [ghead]
    for gi, ids, cats, survivor, (s, n, bd) in group_info:
        years = [int(cards[i]["bd"][:4]) for i in ids if cards[i]["bd"] not in SENTINELS]
        for i in ids:
            r = [gi, "; ".join(sorted(cats)), i, "да" if i == survivor else "", cards[i]["s0"], cards[i]["n0"],
                 cards[i]["bd"], s, n, bd, cards[i]["sex"], cards[i]["nl"], cards[i]["nr"]]
            grows.append(r)
            if years and max(years) - min(years) >= 15:
                gap.append(r)
    gw = [7, 30, 8, 9, 22, 18, 11, 22, 18, 11, 6, 8, 11]
    for name_, rows in (("Склейка", grows), ("Склейка ДР ±15 лет", gap)):
        ws = sheet(name_, rows, gw)
        for k, r in enumerate(ws.iter_rows(min_row=2), 1):
            for c in r:
                c.fill = fill[(r[0].value or 0) % 2]
    unresolved = [i for i, c in cards.items() if c["one_word"] and i not in merged_into and i not in final]
    sheet("Не решено", [head] + [row(i) for i in sorted(unresolved)], widths)
    ws0.append(["Показатель", "Значение"])
    for k, v in (("Карточек всего", len(cards)), ("Групп склейки", len(group_info)),
                 ("Карточек удалится (склеены)", len(merged_into)),
                 ("Карточек после чистки", len(cards) - len(merged_into)),
                 ("Исправлено ФИО без склейки", len(fio)), ("Латиница (перевод/оставлено)", len(lat)),
                 ("Групп с разницей ДР ≥ 15 лет (родитель/ребёнок?)", len({r[0] for r in gap[1:]})),
                 ("Не решено", len(unresolved)), ("Совпадений ФИО+ДР после чистки (должно быть 0)", len(clash))):
        ws0.append([k, v])
    ws0.column_dimensions["A"].width = 55
    ws0.column_dimensions["B"].width = 12
    ws0.append([])
    ws0.append(["Листы «Латиница» и «ФИО»: правьте колонки «Станет фамилия/имя» — при применении берутся из файла."])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--report", required=True, help="xlsx было/станет")
    ap.add_argument("--review", help="проверенный пользователем xlsx (правки «Станет …»)")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--backup")
    args = ap.parse_args()

    from scripts.import_boom_historical import get_connection
    conn = get_connection()
    try:
        clients, leads, results = load(conn)
        names = Names([(r["surname"], r["name"], r["sex"]) for r in leads + results])
        overrides = read_overrides(args.review) if args.review else {}
        cards = build_cards(clients, leads, results, names, overrides)
        final, merged_into, group_info, clash = plan_changes(cards, names)
        write_report(args.report, cards, final, merged_into, group_info, clash)
        print(f"Карточек: {len(cards)}, групп склейки: {len(group_info)}, удалится карточек: {len(merged_into)}, "
              f"изменится ФИО/ДР: {len(final)}, правок из файла: {len(overrides)}, совпадений после чистки: {len(clash)}")
        print(f"Отчёт: {args.report}")
        if not args.apply:
            print("\ndry-run. Повтори с --review <файл> --apply --backup <путь>.")
            return 0
        if not args.backup or clash:
            print("Нужен --backup." if not args.backup else f"Совпадения ФИО+ДР после чистки: {clash[:5]} — не применяю.")
            return 1
        nl, nr = apply(conn, cards, final, merged_into, args.backup)
        print(f"Применено: склеено карточек {len(merged_into)}, изменено {len(final)}; заявок затронуто {nl}, "
              f"результатов {nr}. Бэкап: {args.backup}")
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
