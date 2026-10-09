"""
Импорт архива Siberman 2016–2024 со всеми КТ из таблицы организатора
«Online results Siberman.xlsx» (листы по годам + «Рейтинг …»).

Формат листов — тот, в котором в БД уже лежит 2025: вело-1 — время от
старта гонки (с плаванием), вело-2 — чистое время от своей стартовой
минуты, бег — от старта этапа. Админская загрузка не подходит (там вело —
астрономическое время).

- КТ года заводятся по заголовкам листа: плавание и бег — seq по
  дистанции (на seq завязан счётчик кругов), вело — по порядку, финиш
  этапа всегда STAGE_MAX_SEQ; пустые во всём листе колонки не заводятся.
- Пола в листах нет: личники — из «Рейтинг Мужчины/Женщины», эстафетчики —
  по имени/фамилии. Эстафетчики, записанные одной фамилией (2017–2019), —
  полные ФИ из «Рейтинг Эстафеты» (решение пользователя 2026-10-09).
- Рекорды не пересчитываются (apply_to_db(update_records=False)).

  python scripts/siberman_import_history.py                         # dry-run 2016–2024
  python scripts/siberman_import_history.py --years 2016 2017 --apply
"""

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")
load_dotenv(ROOT / ".env.local", override=True)  # локально — MySQL на 3308, на проде файла нет

import openpyxl

from src.siberman.db import get_siberman_connection
from src.siberman.finish_counts import person_key
from src.siberman.parser import ParseResult, _compute_dnf_stage, parse_time_to_seconds
from src.siberman.service import (
    STAGE_MAX_SEQ, apply_to_db, compute_overall, compute_stage_totals, format_seconds,
)

DEFAULT_XLSX = r"C:\Users\podbo\Downloads\Online results Siberman.xlsx"
ARCHIVE_YEARS = range(2016, 2025)  # 2025+ уже в БД из живой системы — не трогаем

SWIM_SEQ_BY_KM = {1.3: 1, 2.6: 2, 3.9: 3, 5.2: 4, 6.5: 5, 7.8: 6, 10.0: 7}
SWIM_LABELS = {1: "Разворот 1 (1,3 км)", 2: "1 круг (2,6 км)", 3: "Разворот 2 (3,9 км)",
               4: "2 круга (5,2 км)", 5: "Разворот 3 (6,5 км)", 6: "3 круга (7,8 км)", 7: "Финиш (10 км)"}

# Порядок «Фамилия Имя» как в 2024–2026 (в протоколе 2023 переставлены).
NAME_FIXES = {(2023, "Herinckx", "Wijnand"): ("Wijnand", "Herinckx")}

# Официальный итог (протокол + «Рейтинг») расходится с суммой КТ — финишная
# КТ этапа подгоняется под официальный итог (решение пользователя 2026-10-09).
OFFICIAL_FINISH_FIXES = {
    # 28:33:11: сумма двух дней в протоколе 19:01:08 → вело-2 = 19:01:08 − 9:04:30
    (2022, "27"): {("bike_day2", 8): 9 * 3600 + 56 * 60 + 38},
    # «Общее время 84 км» 8:46:53, а 12-й круг — 8:46:38
    (2023, "3"): {("run", 12): 8 * 3600 + 46 * 60 + 53},
}

MALE_NAMES_ENDING_A = {"никита", "илья", "кузьма", "лука", "фома", "савва", "данила", "гаврила"}
RELAY_STAGES = ("swim", "bike", "run")
RELAY_STAGE_OF = {"swim": "swim", "bike_day1": "bike", "bike_day2": "bike", "run": "run"}


def norm(v) -> str:
    return " ".join(str(v or "").split())


def split_full_name(full: str) -> tuple[str, str]:
    parts = norm(full).split(None, 1)
    return (parts[0], parts[1]) if len(parts) == 2 else (norm(full), "")


def guess_gender(surname: str, name: str) -> str:
    """Пол эстафетчика по имени (окончание -а/-я), иначе по фамилии."""
    n, s = name.lower().replace("ё", "е"), surname.lower().replace("ё", "е")
    if n:
        return "F" if n[-1] in "ая" and n not in MALE_NAMES_ENDING_A else "M"
    return "F" if re.search(r"(ова|ева|ина|ына|ая)$", s) else "M"


def km_of(header: str) -> float:
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*км", header)
    if not m:
        raise ValueError(f"нет дистанции в заголовке «{header}»")
    return float(m.group(1).replace(",", "."))


def is_kt_header(header: str) -> bool:
    h = header.lower()
    return "км" in h and not h.startswith(("место", "темп", "время", "общее", "средн"))


def assign_seqs(stage: str, headers: list[str]) -> list[tuple[int, str, float]]:
    """Заголовки КТ этапа по порядку → [(seq, label, km)]."""
    out = []
    if stage == "swim":
        for h in headers:
            seq = SWIM_SEQ_BY_KM[km_of(h)]
            out.append((seq, SWIM_LABELS[seq], km_of(h)))
    elif stage == "run":
        for h in headers:
            lap = int(re.match(r"\s*(\d+)\s*круг", h).group(1))
            label = "Финиш — 12 круг (84 км)" if lap == 12 else f"{lap} круг ({lap * 7} км)"
            out.append((lap, label, km_of(h)))
    else:
        *mid, finish = headers
        if not finish.lower().startswith("финиш") or len(mid) >= STAGE_MAX_SEQ[stage]:
            raise ValueError(f"{stage}: неожиданные КТ {headers}")
        out = [(i, h, km_of(h)) for i, h in enumerate(mid, start=1)]
        out.append((STAGE_MAX_SEQ[stage], f"Финиш ({km_of(finish):g} км)", km_of(finish)))
    return out


def sheet_layout(rows: list[tuple]) -> dict:
    """Колонки листа года: КТ по этапам [(col, seq, label, km)], стартовая
    минута вело-2, итог гонки."""
    h0 = [norm(c) for c in rows[0]]
    h1 = [norm(c) for c in rows[1]]

    def col(marker: str) -> int:
        return next(j for j, c in enumerate(h0) if c.startswith(marker))

    data_rows = [r for r in rows[2:] if norm(r[0]) in ("Лично", "Эстафета")]

    def has_data(j: int) -> bool:
        return any(r[j] not in (None, "") for r in data_rows)

    b1, b1_end, start_minute = col("Велоэтап 1 день"), col("Время велоэтапа"), col("Стартовая минута")
    b2, b2_end, run, total = col("Велоэтап 2 день"), col("Средняя скорость 2"), col("Беговой этап"), col("Сумма трех")
    ranges = {"swim": (6, b1), "bike_day1": (b1, b1_end), "bike_day2": (b2, b2_end), "run": (run, total)}

    stages = {}
    for stage, (a, b) in ranges.items():
        cols = [j for j in range(a, b) if is_kt_header(h1[j]) and has_data(j)
                and (stage != "run" or re.match(r"\d+\s*круг", h1[j]))]
        seqs = assign_seqs(stage, [h1[j] for j in cols])
        stages[stage] = [(j, seq, label, km) for j, (seq, label, km) in zip(cols, seqs)]

    # 2020: в колонках бега время КРУГА, а не от старта этапа — значения не
    # растут от КТ к КТ у большинства (у одного-двух — это опечатки).
    def increasing(r) -> bool:
        vals = [parse_time_to_seconds(r[j]) for j, *_ in stages["run"]]
        vals = [v for v in vals if v]
        return all(a < b for a, b in zip(vals, vals[1:]))
    run_laps = sum(not increasing(r) for r in data_rows) > len(data_rows) / 2
    return {"stages": stages, "start_minute": start_minute, "total": total, "place": total + 1,
            "run_laps": run_laps}


def load_ratings(wb) -> tuple[dict, dict]:
    """(gender_by[(person_key, year)], relay_by_year[year] = [(team, [пловец, вело, бег])])."""
    gender_by = {}
    for sheet, g in (("Рейтинг  Мужчины", "M"), ("Рейтинг  Женщины", "F")):
        for r in wb[sheet].iter_rows(values_only=True):
            if isinstance(r[3], (int, float)):
                gender_by[(person_key(norm(r[1]), ""), int(r[3]))] = g
    relay_by_year: dict[int, list] = {}
    for r in wb["Рейтинг  Эстафеты"].iter_rows(values_only=True):
        if isinstance(r[6], (int, float)):
            relay_by_year.setdefault(int(r[6]), []).append((norm(r[1]), [norm(x) for x in r[2:5]]))
    return gender_by, relay_by_year


def fill_relay_names(members: list[str], candidates: list[tuple[str, list[str]]]) -> list[str]:
    """Члены, записанные одной фамилией, → ФИ из «Рейтинг Эстафеты» той же
    команды (команда — та, где совпали фамилии минимум 2 из 3 этапов;
    у несовпавшей — замена в последний момент, берём рейтинг)."""
    if all(len(m.split()) >= 2 for m in members):
        return members
    sur = lambda full: person_key(full.split()[0], "") if full else ""
    for _, rated in candidates:
        if sum(sur(a) == sur(b) for a, b in zip(members, rated)) >= 2:
            return [rated[i] if len(m.split()) < 2 else m for i, m in enumerate(members)]
    raise ValueError(f"эстафета {members}: нет команды в «Рейтинг Эстафеты»")


def drop_outliers(cp: dict, who: str) -> list[str]:
    """Промежуточная КТ не между соседними (опечатка в протоколе, напр.
    «5:33:43» между 2:23:32 и 4:01:23) → None, в отчёт."""
    notes = []
    for stage in STAGE_MAX_SEQ:
        keys = [k for k in sorted(cp) if k[0] == stage and cp[k] is not None]
        bad = [(prev, cur, nxt) for prev, cur, nxt in zip(keys, keys[1:], keys[2:])
               if cp[prev] < cp[nxt] and not cp[prev] < cp[cur] < cp[nxt]]
        for prev, cur, nxt in bad:
            notes.append(f"{who}, {stage} seq {cur[1]}: выброс {format_seconds(cp[cur])} "
                         f"между {format_seconds(cp[prev])} и {format_seconds(cp[nxt])} — пропущено")
        for _, cur, _ in bad:
            cp[cur] = None
    return notes


def row_status(row: tuple, total_col: int) -> str:
    t = norm(row[total_col]).upper()
    if t in ("DNF", "DSQ"):
        return t.lower()
    return "dnf" if any(norm(c).upper() in ("DNF", "ДНФ") for c in row) else "active"


def build_year(wb, year: int, gender_by: dict, relay_by_year: dict) -> tuple[ParseResult, list, list[str]]:
    """Лист года → (ParseResult, КТ [(stage, seq, label, km)], заметки отчёта)."""
    rows = list(wb[str(year)].iter_rows(values_only=True))
    lay = sheet_layout(rows)
    result = ParseResult(race_year=year)
    notes: list[str] = []

    for r in rows[2:]:
        fmt = norm(r[0])
        if fmt not in ("Лично", "Эстафета"):
            continue
        bib = str(int(r[1])) if isinstance(r[1], float) else norm(r[1])
        who = f"№{bib} {norm(r[2])}"
        cp: dict = {}
        for stage, cols in lay["stages"].items():
            for j, seq, label, _ in cols:
                s = parse_time_to_seconds(r[j])
                if s is None and isinstance(r[j], str) and "::" in r[j]:
                    s = parse_time_to_seconds(r[j].replace("::", ":"))
                    notes.append(f"{who}, {stage} «{label}»: «{r[j]}» → {format_seconds(s)}")
                cp[(stage, seq)] = s if s else None  # 0:00:00 — пустая ячейка формулы
                if s is None and norm(r[j]) and norm(r[j]).upper() not in ("DNF", "DNS", "DSQ"):
                    notes.append(f"{who}, {stage} «{label}»: не распознано «{r[j]}» — пропущено")
        if lay["run_laps"]:
            acc = 0
            for seq in sorted(seq for stage, seq in cp if stage == "run"):
                lap = cp[("run", seq)]
                acc = acc + lap if lap is not None and acc is not None else None
                cp[("run", seq)] = acc
        for key, official in OFFICIAL_FINISH_FIXES.get((year, bib), {}).items():
            notes.append(f"{who}, {key[0]} финиш: {format_seconds(cp[key])} → {format_seconds(official)} (официальный итог)")
            cp[key] = official
        notes += drop_outliers(cp, who)
        status = row_status(r, lay["total"])
        start_minute = parse_time_to_seconds(r[lay["start_minute"]])

        if fmt == "Лично":
            surname, name = norm(r[2]), norm(r[3])
            surname, name = NAME_FIXES.get((year, surname, name), (surname, name))
            gender = gender_by.get((person_key(surname, name), year))
            if gender is None:
                raise ValueError(f"{year} №{bib} {surname} {name}: нет в «Рейтинге» — пол неизвестен")
            result.participants.append({
                "race_year": year, "bib": bib, "surname": surname, "name": name, "gender": gender,
                "country": norm(r[4]) or "Россия", "city": norm(r[5]), "format": "individual",
                "status": status, "dnf_stage": _compute_dnf_stage(cp) if status == "dnf" else None,
                "_cp_key": bib,
            })
            result.checkpoint_times[bib] = cp
            if start_minute:
                result.handicaps[bib] = start_minute
            continue

        members = [norm(m) for m in r[3:6]]
        filled = fill_relay_names(members, relay_by_year.get(year, []))
        if filled != members:
            notes.append(f"эстафета №{bib} «{norm(r[2])}»: {members} → {filled}")
        dnf_stage = _compute_dnf_stage(cp) if status == "dnf" else None
        dnf_role = RELAY_STAGE_OF.get(dnf_stage)
        for role, full in zip(RELAY_STAGES, filled):
            surname, name = split_full_name(full)
            if status == "dnf":
                order = RELAY_STAGES.index
                m_status = "active" if order(role) < order(dnf_role) else "dnf" if role == dnf_role else "dns"
            else:
                m_status = status
            result.participants.append({
                "race_year": year, "bib": bib, "format": "relay", "relay_team_name": norm(r[2]),
                "country": "Россия", "city": "", "surname": surname, "name": name,
                "gender": guess_gender(surname, name), "relay_stage": role, "status": m_status,
                "dnf_stage": dnf_stage if role == dnf_role else None, "_cp_key": f"{bib}:{role}",
            })
            result.checkpoint_times[f"{bib}:{role}"] = {
                k: v for k, v in cp.items() if RELAY_STAGE_OF[k[0]] == role}
        if start_minute:
            result.handicaps[f"{bib}:bike"] = start_minute

    checkpoints = [(stage, seq, label, km) for stage, cols in lay["stages"].items() for _, seq, label, km in cols]
    return result, checkpoints, notes


def verify_against_sheet(wb, year: int, result: ParseResult) -> list[str]:
    """(расхождения итога, расхождения места). Итог личника, посчитанный
    нашей моделью, должен совпасть с «Сумма трех дней» — иначе запись
    отменяется. Место — только предупреждение: оно у нас по итогу, а в
    протоколе 2022 места Шапенко/Ошуркова не соответствуют их итогам
    (рейтинг организатора — по итогу, как у нас)."""
    rows = list(wb[str(year)].iter_rows(values_only=True))
    lay = sheet_layout(rows)
    sheet = {}
    for r in rows[2:]:
        if norm(r[0]) == "Лично":
            bib = str(int(r[1])) if isinstance(r[1], float) else norm(r[1])
            sheet[bib] = (parse_time_to_seconds(r[lay["total"]]), r[lay["place"]])
    problems, places, finishers = [], [], []
    for p in result.participants:
        if p["format"] != "individual":
            continue
        ours = compute_overall(compute_stage_totals(result.checkpoint_times[p["bib"]]))
        theirs, _ = sheet[p["bib"]]
        if p["status"] == "active":
            finishers.append((ours, p["bib"]))
            if ours != theirs:
                problems.append(f"№{p['bib']} {p['surname']}: итог {format_seconds(ours)} ≠ лист {format_seconds(theirs)}")
    for place, (_, bib) in enumerate(sorted(f for f in finishers if f[0] is not None), start=1):
        if sheet[bib][1] != place:
            places.append(f"№{bib}: место {place} ≠ лист {sheet[bib][1]:g}")
    return problems, places


def write_checkpoints(conn, year: int, checkpoints: list) -> None:
    cur = conn.cursor()
    cur.execute("DELETE FROM participants WHERE race_year=%s", (year,))
    cur.execute("DELETE FROM checkpoints WHERE race_year=%s", (year,))
    cur.executemany(
        "INSERT INTO checkpoints (race_year, stage, seq, label, distance_km) VALUES (%s, %s, %s, %s, %s)",
        [(year, stage, seq, label, km) for stage, seq, label, km in checkpoints],
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--years", type=int, nargs="*", default=list(ARCHIVE_YEARS))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    if any(y not in ARCHIVE_YEARS for y in args.years):
        sys.exit(f"Только архив {ARCHIVE_YEARS.start}–{ARCHIVE_YEARS.stop - 1}: 2025+ ведёт живая система.")

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    gender_by, relay_by_year = load_ratings(wb)
    built = {}
    for year in args.years:
        result, checkpoints, notes = build_year(wb, year, gender_by, relay_by_year)
        problems, places = verify_against_sheet(wb, year, result)
        ind = [p for p in result.participants if p["format"] == "individual"]
        rel = [p for p in result.participants if p["format"] == "relay"]
        by_status = {s: sum(p["status"] == s for p in ind) for s in ("active", "dnf", "dsq")}
        laps = " (бег в листе — время кругов, суммируем)" if sheet_layout(
            list(wb[str(year)].iter_rows(values_only=True)))["run_laps"] else ""
        print(f"\n=== {year}{laps}: личники {len(ind)} (финиш {by_status['active']}, DNF {by_status['dnf']}, "
              f"DSQ {by_status['dsq']}), эстафет {len(rel) // 3}")
        for stage in STAGE_MAX_SEQ:
            print(f"  КТ {stage}: " + ", ".join(f"{seq}:{label}" for st, seq, label, _ in checkpoints if st == stage))
        print("  женщины-личники: " + (", ".join(f"{p['surname']} {p['name']}" for p in ind if p["gender"] == "F") or "—"))
        print("  женщины-эстафетчицы: " + (", ".join(f"{p['surname']} {p['name']}" for p in rel if p["gender"] == "F") or "—"))
        for p in result.participants:
            if p["status"] != "active":
                print(f"  {p['status'].upper()}: №{p['bib']} {p['surname']} {p['name']} {p.get('relay_stage') or ''} {p['dnf_stage'] or ''}")
        for n in notes:
            print("  " + n)
        print("  сверка с листом: " + ("OK" if not problems else f"{len(problems)} расхождений"))
        for pr in problems + [f"(предупреждение) {pl}" for pl in places]:
            print("    " + pr)
        built[year] = (result, checkpoints, problems)

    if not args.apply:
        print("\ndry-run. Повтори с --apply.")
        return
    if any(pr for _, _, pr in built.values()):
        sys.exit("\nЕсть расхождения со сверкой — запись отменена.")
    conn = get_siberman_connection()
    if conn is None:
        sys.exit("Нет соединения с БД siberman (DB_* в .env/.env.local).")
    try:
        for year, (result, checkpoints, _) in built.items():
            write_checkpoints(conn, year, checkpoints)
            summary = apply_to_db(result, update_records=False)
            if not summary.get("ok"):
                sys.exit(f"{year}: {summary.get('error')}")
            print(f"{year}: записано участников {summary['participants']}, отметок {summary['checkpoint_times']}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
